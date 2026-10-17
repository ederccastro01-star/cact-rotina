# -*- coding: utf-8 -*-
"""Contrato da IA para a Agenda Paralela: só peça, dias fixados no texto e resumo.

Polo, destino e data NÃO vêm daqui (ver polo.py, destino.py, prazos.py). Chamada por HTTP com
``requests`` (convenção do projeto: sem dependências novas na VM). Modelo em ``cfg["modelo_claude"]``.
"""
from __future__ import annotations

import json
import os
import re
import unicodedata

import openpyxl

PADROES = {
    "peca": "Manifestação", "prazo_dias": None, "prazo_origem": "nenhum", "dias_uteis": True,
    "prazo_comum": False, "ente_publico_intimado": False, "tipo_ato": "notificação", "teor_presente": True,
    "favoravel": "neutro", "data_retorno": None, "resumo": "", "honorarios": False,
}
VOCABULARIO_PADRAO = [
    ("Atenção! Honorários", "pagamento, alvará, levantamento, depósito, honorários", None, "ambos"),
    ("Embargos de Declaração", "sentença, decisão ou acórdão: ED em 5 dias antes do recurso", 5, "ambos"),
    ("Recurso Ordinário", "recurso da sentença", 8, "ambos"),
    ("Agravo de Petição", "recurso na execução", 8, "ambos"),
    ("Contrarrazões", "resposta ao recurso ordinário da parte contrária", 8, "ambos"),
    ("Contraminuta ao Agravo de Petição", "resposta ao agravo de petição da parte contrária", 8, "ambos"),
    ("Impugnação aos cálculos", "art. 879 §2º, prazo comum", 8, "ambos"),
    ("Manifestação s/ impugnação aos cálculos", "vista da impugnação da parte contrária", 8, "ambos"),
    ("Manifestação s/ Embargos à Execução", "vista dos embargos (art. 884)", 5, "ambos"),
    ("Manifestação s/ emenda à inicial", "vista da emenda", 5, "ambos"),
    ("Manifestação", "manifestar-se sobre petição, laudo, documento", None, "ambos"),
    ("Informar dados bancários (alvará)", "indicar conta para expedição de alvará", 5, "ambos"),
    ("Requerer prosseguimento da execução", "indicar meios, requerer o que entender de direito", None, "ambos"),
    ("Ciência (sem providência)", "ofício, distribuição, recebimento de recurso, certidão", None, "ambos"),
    ("Acompanhamento: <motivo>", "sobrestamento/suspensão: data de retorno", None, "ambos"),
]
KW_HONORARIOS = ["pagamento", "liberação", "liberacao", "levantamento", "alvará", "alvara", "transferência",
                 "transferencia", "devolução de saldo", "certidão de crédito", "habilitação de crédito",
                 "honorários", "honorarios", "depósito judicial", "deposito judicial"]
KW_DECISAO = ["sentença", "sentenca", "acórdão", "acordao", "decisão", "decisao", "homologo os cálculos", "dispositivo"]
TIPOS_ATO = [("acórdão", r"ac[oó]rd[aã]o"), ("sentença", r"senten[cç]a"), ("decisão", r"\bdecis[aã]o\b"),
             ("despacho", r"\bdespacho\b"), ("edital", r"\bedital\b"), ("distribuição", r"distribu[ií]do"),
             ("pauta", r"\bpauta\b|audi[êe]ncia"), ("certidão", r"certid[aã]o"), ("notificação", r".")]
LIMITE_RESUMO = 350

PROMPT_PARALELA = """Você é assistente jurídico do CACT (Castro Advocacia & Cálculos Trabalhistas), escritório trabalhista de Goiás. Recebe o texto de uma intimação publicada no DJEN e devolve SÓ um JSON. Polo, destino e data do prazo são calculados por código fora daqui: NÃO decida de quem é o prazo e NÃO calcule datas.

ENTRADAS: (1) polo do intimado (AUTOR, RÉU, AUTOR e RÉU ou desconhecido); (2) lado que o escritório representa (AUTOR, RÉU ou desconhecido); (3) vocabulário de peças da Agenda, uma por linha "rótulo | significado | prazo | lado"; (4) texto da intimação.

REGRAS:
1. "peca": a providência cabível AO INTIMADO, usando um rótulo do vocabulário; se nenhum servir, "Outro: <texto curto>". Se o intimado é a parte contrária ao escritório, termine o rótulo com " - aguardar". Se a publicação envolve pagamento, liberação, levantamento, alvará, transferência de valores, devolução de saldo, certidão/habilitação de crédito, depósito judicial ou honorários, comece com "Atenção! Honorários - " e marque "honorarios": true. Se o ato é sentença, decisão ou acórdão, a peça é "Embargos de Declaração" (regra do escritório: 5 dias, antes do recurso), salvo prazo menor fixado no texto; mencione o recurso cabível no resumo. Ato de mera ciência (ofício expedido, distribuição, recebimento de recurso, certidão) → "Ciência (<o que>)".
2. "prazo_dias": o número de dias que o TEXTO fixa (inteiro), ou null se o texto não fixa ("prazo legal", "prazo e fins legais" = null). Nunca invente. "prazo_origem": "texto" quando veio do texto, "nenhum" caso contrário. "dias_uteis": false só se o texto disser dias corridos. "48 horas" → prazo_dias 2.
3. "prazo_comum": true se o texto diz "prazo comum", "ambas as partes", ou intima nos termos do art. 879 §2º da CLT.
4. "ente_publico_intimado": true se o intimado é União, Estado, Município, autarquia ou fundação pública.
5. "tipo_ato": um de sentença, acórdão, decisão, despacho, notificação, edital, distribuição, pauta, certidão.
6. "teor_presente": false quando a intimação só manda tomar ciência de um documento ("despacho ID …") sem transcrevê-lo.
7. "favoravel": para sentença/decisão/acórdão, "sim", "não" ou "neutro" do ponto de vista do lado que o escritório representa; "neutro" se desconhecido.
8. "data_retorno": "AAAA-MM-DD" quando o ato sobresta ou suspende o processo até uma data (recuperação judicial, stay period, sobrestamento): o dia útil seguinte ao fim; senão null.
9. "resumo": UMA linha, até 350 caracteres, neste formato exato e nesta ordem, omitindo partes sem conteúdo: "Intimado: NOME (Autor/Réu) · Ato: <tipo> ID <x> de dd/mm · Determinação/Resultado: <uma frase> · Prazo: <n dias (texto|legal)> · Para nós: <uma frase, se couber> · Valores/honorários: <se houver>". Não repita o cabeçalho do processo nem os nomes dos advogados.

Responda SOMENTE com o JSON, sem texto antes ou depois:
{"peca": "...", "prazo_dias": null|int, "prazo_origem": "texto|nenhum", "dias_uteis": true|false, "prazo_comum": true|false, "ente_publico_intimado": true|false, "tipo_ato": "...", "teor_presente": true|false, "favoravel": "sim|não|neutro", "data_retorno": null|"AAAA-MM-DD", "resumo": "...", "honorarios": true|false}"""

POLO_TXT = {"A": "Autor", "P": "Réu"}


def _asc(s) -> str:
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()


# ----------------------------------------------------------------------------- vocabulário
def carregar_vocabulario(cfg: dict) -> list[dict]:
    caminho = cfg.get("arquivo_cadastro")
    aba = cfg.get("aba_vocabulario", "Vocabulário")
    out: list[dict] = []
    if caminho and os.path.exists(caminho):
        try:
            wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
            if aba in wb.sheetnames:
                for i, row in enumerate(wb[aba].iter_rows(values_only=True)):
                    if i == 0 or not row or not row[0]:
                        continue
                    prazo = row[2] if len(row) > 2 else None
                    try:
                        prazo = int(prazo) if prazo not in (None, "") else None
                    except (TypeError, ValueError):
                        prazo = None
                    out.append({"rotulo": str(row[0]).strip(), "significado": str(row[1] or "").strip() if len(row) > 1 else "",
                                "prazo_dias": prazo, "lado": str(row[3] or "ambos").strip().lower() if len(row) > 3 else "ambos"})
            wb.close()
        except Exception:  # noqa: BLE001 — cadastro ilegível: usa o padrão
            out = []
    if not out:
        out = [{"rotulo": r, "significado": s, "prazo_dias": p, "lado": l} for r, s, p, l in VOCABULARIO_PADRAO]
    return out


# ----------------------------------------------------------------------------- parse
def parse_resposta(conteudo: str) -> dict:
    d = dict(PADROES)
    m = re.search(r"\{.*\}", conteudo or "", re.S)
    if not m:
        return d
    try:
        dados = json.loads(m.group(0))
    except json.JSONDecodeError:
        return d
    if not isinstance(dados, dict):
        return d
    for k in PADROES:
        if k in dados and dados[k] is not None:
            d[k] = dados[k]
    try:
        d["prazo_dias"] = int(d["prazo_dias"]) if d["prazo_dias"] not in (None, "") else None
    except (TypeError, ValueError):
        d["prazo_dias"] = None
    for k in ("dias_uteis", "prazo_comum", "ente_publico_intimado", "teor_presente", "honorarios"):
        d[k] = bool(d[k])
    for k in ("peca", "prazo_origem", "tipo_ato", "favoravel", "resumo"):
        d[k] = str(d[k]).strip() if d[k] is not None else PADROES[k]
    if d["data_retorno"] is not None:
        d["data_retorno"] = str(d["data_retorno"])[:10]
    if not d["peca"]:
        d["peca"] = PADROES["peca"]
    return d


# ----------------------------------------------------------------------------- fallback (sem IA)
def fallback(texto: str, ato) -> dict:
    t = _asc(texto)
    honor = any(_asc(k) in t for k in KW_HONORARIOS)
    decisao = any(_asc(k) in t for k in KW_DECISAO)
    tipo = next(nome for nome, rx in TIPOS_ATO if re.search(rx, t))
    if decisao:
        peca = "Embargos de Declaração"
    elif re.search(r"\bdistribu[ií]do\b|expe[cç]a-se of[ií]cio|recebo o?s? recursos?", t):
        peca = "Ciência (sem providência)"
    else:
        peca = "Manifestação"
    if honor:
        peca = "Atenção! Honorários - " + peca
    polos = sorted(getattr(ato, "polos", set()) or [])
    polo_txt = " e ".join(POLO_TXT.get(p, p) for p in polos) or "?"
    nomes = ", ".join(getattr(ato, "intimados", []) or []) or "não identificado"
    corpo = re.sub(r"\s+", " ", texto or "").strip()
    resumo = f"Intimado: {nomes} ({polo_txt}) · Ato: {tipo} · {corpo[:220]}"
    return dict(PADROES, peca=peca, tipo_ato=tipo, honorarios=honor, resumo=resumo[:LIMITE_RESUMO])


def formatar_resumo(d: dict, intimados: list[str], polo: str) -> str:
    resumo = (d.get("resumo") or "").strip()
    if resumo.startswith("Intimado"):
        return resumo[:LIMITE_RESUMO]
    polo_txt = {"A": "Autor", "P": "Réu", "AP": "Autor e Réu"}.get(polo, polo or "?")
    nomes = ", ".join(intimados) if intimados else "não identificado"
    cab = f"Intimado: {nomes} ({polo_txt}) · Ato: {d.get('tipo_ato') or 'ato'}"
    return (cab + (" · " + resumo if resumo else ""))[:LIMITE_RESUMO]


# ----------------------------------------------------------------------------- chamada
def _post(url, headers=None, json=None, timeout=None):
    import requests  # noqa: WPS433 — opcional no PC; obrigatório na VM
    return requests.post(url, headers=headers, json=json, timeout=timeout)


def _entrada(texto: str, polo_intimado: str, lado_nosso: str, vocabulario: list[dict]) -> str:
    voc = "\n".join(f"{v['rotulo']} | {v.get('significado', '')} | {v.get('prazo_dias') or '-'} | {v.get('lado', 'ambos')}"
                    for v in vocabulario)
    return (f"Polo do intimado: {polo_intimado or 'desconhecido'}\n"
            f"Lado que o escritório representa: {lado_nosso or 'desconhecido'}\n\n"
            f"Vocabulário de peças:\n{voc}\n\n"
            f"Texto da intimação:\n{(texto or '')[:9000]}")


def classificar(texto: str, polo_intimado: str, lado_nosso: str, vocabulario: list[dict], cfg: dict) -> dict | None:
    chave = (cfg.get("anthropic_api_key") or "").strip()
    if not chave:
        return None
    try:
        from cact_rotina import com_tentativas, log
    except ImportError:  # pragma: no cover
        com_tentativas = lambda f, _d: f()  # noqa: E731
        log = print
    corpo = {
        "model": cfg.get("modelo_claude", "claude-opus-5-5"),
        "max_tokens": 1500,
        "output_config": {"effort": "medium"},
        "system": PROMPT_PARALELA,
        "messages": [{"role": "user", "content": _entrada(texto, polo_intimado, lado_nosso, vocabulario)}],
    }
    headers = {"x-api-key": chave, "anthropic-version": "2023-06-01", "content-type": "application/json"}
    try:
        r = com_tentativas(lambda: _post("https://api.anthropic.com/v1/messages", headers=headers, json=corpo,
                                         timeout=cfg.get("timeout_segundos", 60)), "Claude (Paralela)")
        if r.status_code != 200:
            log(f"  ! Claude (Paralela) HTTP {r.status_code}: {str(getattr(r, 'text', ''))[:160]}")
            return None
        dados = r.json()
        if dados.get("stop_reason") == "refusal":
            log("  ! Claude (Paralela) recusou a classificação; usando fallback.")
            return None
        textos = [b.get("text", "") for b in dados.get("content", []) if b.get("type") == "text"]
        return parse_resposta("\n".join(textos))
    except Exception as e:  # noqa: BLE001
        log(f"  ! Claude (Paralela) erro: {e}")
        return None
