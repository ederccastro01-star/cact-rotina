# -*- coding: utf-8 -*-
"""Escrita do arquivo da Agenda Paralela no formato visual da Agenda oficial (A–K) + colunas de controle."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from cact_rotina import limpar_texto

COLUNAS_AGENDA = ["Prazo", "Dia/Sem", "Peça", "Publicado", "Processo ", "Reclamante", "Reclamado", "Advogado ",
                  "Prioridade ", "Cliente", "Despacho"]
COLUNAS_CONTROLE = ["Destino", "Decisão", "Motivo CONFERIR", "Polo intimado", "Intimado(s)", "Fonte do polo",
                    "Prazo no texto", "Prazo legal", "Base legal", "Link (teor)", "ID DJEN", "Lote", "Observações",
                    "Prazo calculado"]
LARGURAS = [10.9, 10.1, 24, 11.4, 26, 30, 34, 26, 12.4, 12, 60, 20, 18, 38, 12, 34, 16, 10, 10, 13, 40, 22, 22, 30, 12]
DECISOES = '"Agendar,Descartar,Corrigir e agendar"'
ORDEM_DESTINO = {"PROCESSO NOVO": 0, "PRAZO DO ESCRITÓRIO": 2, "DISPENSADA": 4, "PARTE CONTRÁRIA": 3, "FORA DA AGENDA": 5}

ARIAL = Font(name="Arial", size=10)
COURIER = Font(name="Courier New", size=10)
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="0070C0")
CTRL_FILL = PatternFill("solid", fgColor="595959")
FILL = {"PRAZO DO ESCRITÓRIO": "E2EFDA", "PARTE CONTRÁRIA": "F2F2F2", "PROCESSO NOVO": "FFF2CC", "CONFERIR": "FCE4D6",
        "DISPENSADA": "F2F2F2", "FORA DA AGENDA": "FFFFFF"}
AMARELO = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BORDA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ESQ = Alignment(horizontal="left", vertical="top", wrap_text=True)
COLS_ESQ = {6, 7, 8, 11, 14, 16, 21, 24}

LEIA_ME = [
    ("O que é este arquivo", True),
    ("Agenda Paralela do dia, gerada às 5h30 pela rotina do CACT a partir das publicações do DJEN. Toda publicação de processo "
     "do escritório virou uma linha; nada foi gravado na Agenda oficial. As colunas A a K têm o formato da Agenda; de L em "
     "diante são controle.", False),
    ("", False),
    ("Como usar, de manhã", True),
    ("1. Aba Lançamentos: amarelo = PROCESSO NOVO; laranja = CONFERIR (há dúvida registrada na coluna N); verde = PRAZO DO "
     "ESCRITÓRIO; cinza = PARTE CONTRÁRIA, só para conhecimento.", False),
    ("2. Preencha a coluna M (Decisão): Agendar (A), Descartar (D) ou Corrigir e agendar (C). Em 'Corrigir e agendar', "
     "ajuste antes a data (A), a peça (C) ou a nota (K) na própria linha.", False),
    ("3. Salve. O script do PC grava as linhas 'Agendar' e 'Corrigir e agendar' na aba Publicações da Agenda oficial, "
     "copiando A a K como estão, e carimba a coluna W (Lote) com 'Transportado em ...'.", False),
    ("4. Linha sem decisão não se perde: amanhã aparece na aba Pendentes e no aviso da manhã.", False),
    ("", False),
    ("O que a máquina decide e o que você decide", True),
    ("Máquina: polo do intimado (DJEN), destino (polo x coluna J), data (dias úteis, feriados, tabela legal; na dúvida a "
     "menor), agrupamento de comunicações repetidas num só ato, conciliação (toda publicação tem linha). IA: só o rótulo "
     "da peça (C) e o resumo (K). Você: a decisão e qualquer ajuste.", False),
    ("", False),
    ("Colunas de controle", True),
    ("L Destino · M Decisão (você) · N Motivo CONFERIR · O Polo intimado · P Intimado(s) · Q Fonte do polo · R Prazo no "
     "texto · S Prazo legal · T Base legal · U Link para o teor · V ID DJEN · W Lote/Transportado · X Observações.", False),
]


@dataclass
class Linha:
    prazo: date | None
    peca: str
    publicado: date
    processo: str
    reclamante: str
    reclamado: str
    advogado: str
    prioridade: str
    cliente: str
    despacho: str
    destino: str
    motivo: str = ""
    polo_intimado: str = "?"
    intimados: str = ""
    fonte_polo: str = ""
    prazo_texto: int | None = None
    prazo_legal: int | None = None
    base_legal: str | None = None
    link: str = ""
    ids: str = ""
    lote: str = ""
    observacoes: str = ""


def _ordem(l: Linha):
    if l.destino == "PROCESSO NOVO":
        grupo = 0
    elif l.motivo and l.destino == "PRAZO DO ESCRITÓRIO":
        grupo = 1
    else:
        grupo = ORDEM_DESTINO.get(l.destino, 9)
    return (grupo, l.prazo or date.max, l.processo)


def _cabecalho(ws, titulo: str, subtitulo: str):
    ws["A1"] = titulo; ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"] = subtitulo; ws["A2"].font = Font(name="Arial", size=9, italic=True, color="595959")
    ws["A4"] = ("Colunas A a K: iguais à Agenda oficial. Colunas a partir de L: controle da Paralela "
                "(não vão para a Agenda, exceto ID DJEN, Link e Lote).")
    ws["A4"].font = Font(name="Arial", size=9, color="595959")
    ws["A5"] = ("Legenda: verde = PRAZO DO ESCRITÓRIO; cinza = PARTE CONTRÁRIA; amarelo = PROCESSO NOVO; laranja = CONFERIR. "
                "Você preenche só a coluna M (Decisão) e, se quiser, A/C/K e X (Observações).")
    ws["A5"].font = Font(name="Arial", size=9, color="595959")
    for i, (h, w) in enumerate(zip(COLUNAS_AGENDA + COLUNAS_CONTROLE, LARGURAS), start=1):
        c = ws.cell(row=7, column=i, value=h)
        c.font = HDR_FONT; c.fill = HDR_FILL if i <= 11 else CTRL_FILL; c.alignment = CENTRO; c.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[7].height = 30
    ws.freeze_panes = "F8"


def _escrever(ws, row: int, l: Linha):
    cor = FILL["CONFERIR"] if (l.motivo and l.destino == "PRAZO DO ESCRITÓRIO") else FILL.get(l.destino, "FFFFFF")
    fill = PatternFill("solid", fgColor=cor)
    vals = [l.prazo, None, l.peca, l.publicado, l.processo, l.reclamante, l.reclamado, l.advogado, l.prioridade, l.cliente,
            l.despacho, l.destino, None, l.motivo, l.polo_intimado, l.intimados, l.fonte_polo, l.prazo_texto, l.prazo_legal,
            l.base_legal, l.link, l.ids, l.lote, l.observacoes, l.prazo]       # Y = prazo calculado (fica mesmo se A for corrigida)
    for i, v in enumerate(vals, start=1):
        if isinstance(v, str):
            v = limpar_texto(v)
        c = ws.cell(row=row, column=i, value=v)
        c.font = COURIER if i == 5 else ARIAL; c.border = BORDA; c.fill = fill
        c.alignment = ESQ if i in COLS_ESQ else CENTRO
        if i in (1, 4, 25):
            c.number_format = "dd/mm/yy;@"
        if i == 2:
            c.value = f'=IF(A{row}="","",A{row})'; c.number_format = "ddd"
        if i == 5:
            c.number_format = "@"
        if i == 21 and v:
            c.hyperlink = v; c.value = "abrir teor"
            c.font = Font(name="Arial", size=9, color="0563C1", underline="single")
    ws.cell(row=row, column=13).fill = AMARELO
    ws.row_dimensions[row].height = 54


def _aba(wb, titulo, subtitulo, nome, linhas: list[Linha], ordenar=True):
    ws = wb.create_sheet(nome)
    _cabecalho(ws, titulo, subtitulo)
    row = 8
    for l in (sorted(linhas, key=_ordem) if ordenar else linhas):
        _escrever(ws, row, l); row += 1
    if linhas:
        dv = DataValidation(type="list", formula1=DECISOES, allow_blank=True)
        dv.prompt = ("Agendar: o script grava na Agenda oficial. Descartar: só registro. "
                     "Corrigir e agendar: corrija A, C ou K nesta linha antes.")
        ws.add_data_validation(dv); dv.add(f"M8:M{row - 1}")
    return ws, row


def gerar_paralela(lancamentos: list[Linha], fora: list[Linha], caminho: str, data_disp: date, conciliacao: dict) -> str:
    wb = openpyxl.Workbook()
    wl = wb.active; wl.title = "Leia-me"
    for i, (t, b) in enumerate(LEIA_ME, start=1):
        c = wl.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", size=11 if b else 10, bold=b); c.alignment = Alignment(wrap_text=True, vertical="top")
    wl.column_dimensions["A"].width = 140

    ws, row = _aba(wb, f"CACT - Agenda Paralela {data_disp:%d-%m-%Y}",
                   "Gerada às 5h30 pela rotina do CACT. Ordem: PROCESSO NOVO e CONFERIR no topo; depois PRAZO DO ESCRITÓRIO "
                   "por data; PARTE CONTRÁRIA por último. Polo e data por código; peça e resumo pela IA.",
                   "Lançamentos", lancamentos)
    n_esc = sum(1 for l in lancamentos if l.destino == "PRAZO DO ESCRITÓRIO")
    n_con = sum(1 for l in lancamentos if l.destino == "PARTE CONTRÁRIA")
    n_novo = sum(1 for l in lancamentos if l.destino == "PROCESSO NOVO")
    n_conf = sum(1 for l in lancamentos if l.motivo and l.destino == "PRAZO DO ESCRITÓRIO")
    ws.cell(row=row + 1, column=1, value=(
        f"Conciliação: {conciliacao.get('atos', len(lancamentos) + len(fora))} atos de {conciliacao.get('processos', '?')} processos, "
        f"{conciliacao.get('ids', '?')} comunicações; {n_esc} prazo do escritório, {n_con} parte contrária, {n_novo} processo novo, "
        f"{n_conf} CONFERIR; {len(fora)} fora da Agenda.")).font = Font(name="Arial", size=9, italic=True)
    if not conciliacao.get("ok", True):
        faltam = ", ".join(str(x) for x in conciliacao.get("faltam", []))[:300]
        c = ws["A6"]
        c.value = f"ATENÇÃO: conciliação falhou — comunicações sem linha: {faltam}"
        c.font = Font(name="Arial", size=10, bold=True, color="C00000")

    _aba(wb, "Fora da Agenda — carteira dos parceiros em processos em que o escritório ainda não atuou",
         "Só para consulta; sem decisão obrigatória. Se um deles virar trabalho, lance na Agenda como hoje e a partir daí "
         "ele entra em Lançamentos.", "Fora da Agenda", fora)
    _aba(wb, "Pendentes — linhas de dias anteriores sem decisão",
         "Preenchida pelo script do PC a partir das Paralelas anteriores. Some quando a decisão é dada.", "Pendentes", [])
    wb.save(caminho)
    return caminho
