# -*- coding: utf-8 -*-
import os
from datetime import date

import openpyxl
import pytest

from paralela import gerar
from paralela.destino import RefProc

PROC1 = "00000011120265180001"
COL = {"prazo": 1, "peca": 3, "proc": 5, "cliente": 10, "despacho": 11, "destino": 12, "motivo": 14, "polo": 15, "ids": 22}


def _ref_sintetico(_cfg):
    return {PROC1: RefProc(ref={"Reclamante": "AUTOR TESTE", "Reclamado": "REU TESTE LTDA", "Advogado": "EDER",
                                "Prioridade": "NORMAL", "Cliente": "AUTOR(A)", "Despacho": "nota antiga"},
                           js={"AUTOR(A)"}, conferido=True, inPub=True)}


def _linhas(ws):
    return [r for r in ws.iter_rows(min_row=8, values_only=True) if r and r[4]]


def test_gerar_sintetico_sem_ia(sintetico, cfg, tmp_path, monkeypatch):
    monkeypatch.setattr(gerar, "indexar_agenda_completa", _ref_sintetico)
    cfg["pasta_saida"] = str(tmp_path); cfg["anthropic_api_key"] = ""
    caminho = gerar.gerar_agenda_paralela(sintetico, cfg, date(2026, 10, 16), usar_ia=False)
    assert os.path.basename(caminho) == "CACT - Agenda Paralela 16-10-2026.xlsx"
    wb = openpyxl.load_workbook(caminho)
    lanc = _linhas(wb["Lançamentos"]); fora = _linhas(wb["Fora da Agenda"])
    assert len(lanc) == 4 and len(fora) == 1
    por_ids = {r[COL["ids"] - 1]: r for r in lanc}
    a1 = por_ids["9001;9004"]
    assert a1[COL["destino"] - 1] == "PRAZO DO ESCRITÓRIO"
    assert a1[COL["prazo"] - 1].date() == date(2026, 10, 26)          # 5 dias do texto a partir de 19/10
    assert a1[COL["peca"] - 1].startswith("Atenção! Honorários - ")
    assert a1[COL["cliente"] - 1] == "AUTOR(A)" and a1[COL["despacho"] - 1].startswith("Intimado: AUTOR TESTE")
    a2 = por_ids["9002"]
    assert a2[COL["destino"] - 1] == "PARTE CONTRÁRIA" and a2[COL["peca"] - 1].endswith(" - aguardar")
    a6 = por_ids["9006"]
    assert a6[COL["destino"] - 1] == "PRAZO DO ESCRITÓRIO" and "prazo comum" in a6[COL["motivo"] - 1]
    assert a6[COL["prazo"] - 1].date() == date(2026, 10, 29)          # 8 dias (texto e CLT 879 §2º)
    a3 = por_ids["9003"]
    assert a3[COL["peca"] - 1] == "Embargos de Declaração" and a3[COL["prazo"] - 1].date() == date(2026, 10, 26)
    f = fora[0]
    assert f[COL["destino"] - 1] == "FORA DA AGENDA" and f[COL["ids"] - 1] == "9005" and f[COL["polo"] - 1] == "AUTOR"


def test_conciliacao(sintetico, cfg, tmp_path, monkeypatch):
    monkeypatch.setattr(gerar, "indexar_agenda_completa", _ref_sintetico)
    cfg["pasta_saida"] = str(tmp_path)
    caminho = gerar.gerar_agenda_paralela(sintetico, cfg, date(2026, 10, 16), usar_ia=False)
    wb = openpyxl.load_workbook(caminho)
    ids = []
    for aba in ("Lançamentos", "Fora da Agenda"):
        for r in _linhas(wb[aba]):
            ids += str(r[COL["ids"] - 1]).split(";")
    assert sorted(ids) == sorted(str(p["item"]["id"]) for p in sintetico)
    assert "ATENÇÃO" not in str(wb["Lançamentos"]["A6"].value or "")


def test_data_retorno_iso_invalida_nao_quebra(sintetico, cfg, tmp_path, monkeypatch):
    monkeypatch.setattr(gerar, "indexar_agenda_completa", _ref_sintetico)
    monkeypatch.setattr(gerar, "classificar", lambda *a, **k: {"peca": "Acompanhamento: sobrestado", "prazo_dias": None,
                        "prazo_origem": "nenhum", "dias_uteis": True, "prazo_comum": False, "ente_publico_intimado": False,
                        "tipo_ato": "decisão", "teor_presente": True, "favoravel": "neutro", "data_retorno": "amanhã",
                        "resumo": "Intimado: X (Autor) · Ato: decisão", "honorarios": False})
    cfg["pasta_saida"] = str(tmp_path); cfg["anthropic_api_key"] = "sk"
    caminho = gerar.gerar_agenda_paralela(sintetico[:1], cfg, date(2026, 10, 16), usar_ia=True)
    r = _linhas(openpyxl.load_workbook(caminho)["Lançamentos"])[0]
    assert r[COL["prazo"] - 1].date() == date(2026, 10, 26)        # 5 dias do texto; a data inválida não virou retorno
    assert "data de retorno inválida" in r[COL["motivo"] - 1]


def test_reais_gera_sem_erro(reais, cfg, agenda_amostra, tmp_path):
    cfg["pasta_saida"] = str(tmp_path); cfg["arquivo_agenda"] = agenda_amostra; cfg["anthropic_api_key"] = ""
    caminho = gerar.gerar_agenda_paralela(reais, cfg, date(2026, 10, 16), usar_ia=False)
    wb = openpyxl.load_workbook(caminho)
    lanc = _linhas(wb["Lançamentos"])
    assert len(lanc) >= 30
    assert any(r[COL["destino"] - 1] == "PROCESSO NOVO" and "0011757-98.2024.5.18.0009" in str(r[COL["proc"] - 1]) for r in lanc)
    assert len(_linhas(wb["Fora da Agenda"])) >= 50
    assert "ATENÇÃO" not in str(wb["Lançamentos"]["A6"].value or "")


def test_precatorio_de_segundo_grau_vai_para_lancamentos_ligado_a_origem(cfg, tmp_path, monkeypatch):
    """Pedido do advogado (18/10/2026): processo iniciado no 2º grau (final .0000) sem linha na Agenda e sem OAB nossa
    na comunicação não pode cair em 'Fora da Agenda'. Se o texto cita o processo de origem que está na Agenda, a linha
    herda o lado (J) da origem e vai para Lançamentos com CONFERIR explicando o vínculo."""
    import json
    import cact_rotina as cr
    item = {"id": 7001, "hash": "h7001", "siglaTribunal": "TRT18", "tipoComunicacao": "Intimação", "tipoDocumento": "Intimação",
            "data_disponibilizacao": "2026-10-16", "numero_processo": "00121473220235180000", "numeroprocessocommascara": "0012147-32.2023.5.18.0000",
            "link": "https://pje.trt18.jus.br/pjekz/validacao/7001",
            "texto": "PODER JUDICIÁRIO TRT 18ª REGIÃO OJC RPV E PRECATÓRIOS Precat 0012147-32.2023.5.18.0000 REQUERENTE: AUTOR TESTE "
                     "REQUERIDO: REU TESTE LTDA INTIMAÇÃO Fica V. Sa. intimado para manifestar sobre a atualização do precatório, "
                     "autos de origem 0000001-11.2026.5.18.0001, no prazo de 5 dias.<br>Intimado(s) / Citado(s)<br> - AUTOR TESTE<br>",
            "destinatarios": [{"nome": "AUTOR TESTE", "polo": "A"}], "destinatarioadvogados": []}
    arq = tmp_path / "precat.json"; arq.write_text(json.dumps([item], ensure_ascii=False), encoding="utf-8")
    pubs = cr.carregar_de_json(str(arq))
    monkeypatch.setattr(gerar, "indexar_agenda_completa", _ref_sintetico)
    cfg["pasta_saida"] = str(tmp_path); cfg["anthropic_api_key"] = ""
    wb = openpyxl.load_workbook(gerar.gerar_agenda_paralela(pubs, cfg, date(2026, 10, 16), usar_ia=False))
    lanc = _linhas(wb["Lançamentos"]); fora = _linhas(wb["Fora da Agenda"])
    assert len(lanc) == 1 and len(fora) == 0
    l = lanc[0]
    assert l[COL["proc"] - 1] == "0012147-32.2023.5.18.0000" and l[COL["destino"] - 1] == "PRAZO DO ESCRITÓRIO"
    assert l[COL["cliente"] - 1] == "AUTOR(A)" and "0000001-11.2026.5.18.0001" in l[COL["motivo"] - 1] and "2º grau" in l[COL["motivo"] - 1]
