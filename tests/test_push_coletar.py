# -*- coding: utf-8 -*-
"""Orquestração da coleta do Push no PC (spec §9c), sem Outlook: a listagem é injetada."""
import json
import os
from datetime import date, datetime

import openpyxl

from pc import push_coletar as pc_
from pc.push import caminho_push
from pc.push_fontes import Email
from tests.test_push_fontes import JUS_ANTIGO, PJE


def _cfg(cfg, tmp_path, **push):
    cfg["pc"] = {"pasta_controladoria": str(tmp_path), "hora_alerta_sem_arquivo": "10:00", "intervalo_min": 15}
    cfg["subpasta_push"] = "CACT - PUSH Diário"
    os.makedirs(tmp_path / "CACT - PUSH Diário", exist_ok=True)
    cfg["push"] = {"fontes": [{"conta": "x@gmail.com", "pastas": ["Caixa de Entrada"]}], "janela_dias": 7, "arquivar_em": None,
                   "ledger": str(tmp_path / "push_ledger.json")}
    cfg["push"].update(push)
    return cfg


EMAILS = [
    Email("ederccastro@hotmail.com", "ENC: [TRT18] [PUSH] Atualizações ... 0000001-11.2026.5.18.0011", PJE, datetime(2026, 10, 16, 8, 52), "<pje1@trt18>", "x@gmail.com", "Caixa de Entrada"),
    Email("publicacoes-diarios@jusbrasil.com.br", "2 novas publicações encontradas para você", JUS_ANTIGO, datetime(2026, 10, 16, 10, 2), "<jus1@jusbrasil>", "x@gmail.com", "Caixa de Entrada"),
    Email("nao-responda@jusbrasil.com.br", "Conheça o Jus IA", "promo", datetime(2026, 10, 16, 11, 0), "<mkt@jusbrasil>", "x@gmail.com", "Caixa de Entrada"),
]
AGORA = datetime(2026, 10, 19, 8, 45)


def _listar(chamadas=None):
    def listar(fontes, desde):
        if chamadas is not None:
            chamadas.append((fontes, desde))
        return list(EMAILS)
    return listar


def test_primeira_coleta_gera_relatorio_e_razao(cfg, tmp_path):
    _cfg(cfg, tmp_path)
    chamadas = []
    info = pc_.coletar(cfg, date(2026, 10, 19), AGORA, listar=_listar(chamadas), contar_backlog=lambda fontes, antes: 2)
    assert chamadas[0][1] == datetime(2026, 10, 12, 8, 45) and chamadas[0][0][0]["conta"] == "x@gmail.com"
    assert info["elegiveis"] == 2 and info["novos"] == 2 and info["registros"] == 5 and info["backlog"] == 2        # 3 eventos PJe + 2 Jusbrasil
    assert os.path.basename(info["caminho"]) == "Relatorio_PUSH_2026-10-19_0845.xlsx"
    assert caminho_push(cfg, date(2026, 10, 19)) == info["caminho"]
    ws = openpyxl.load_workbook(info["caminho"])[pc_.ABA]
    assert ws.max_row == 6 and {r[12] for r in ws.iter_rows(min_row=2, values_only=True)} == {"PJe Push", "Jusbrasil"}
    razao = json.load(open(cfg["push"]["ledger"], encoding="utf-8"))
    assert set(razao) == {"<pje1@trt18>", "<jus1@jusbrasil>"} and razao["<pje1@trt18>"]["registros"] == 3
    assert razao["<pje1@trt18>"]["relatorio"] == "Relatorio_PUSH_2026-10-19_0845.xlsx"


def test_segunda_coleta_no_mesmo_dia_nao_repete_nem_gera(cfg, tmp_path):
    _cfg(cfg, tmp_path)
    pc_.coletar(cfg, date(2026, 10, 19), AGORA, listar=_listar(), contar_backlog=lambda f, a: 0)
    info = pc_.coletar(cfg, date(2026, 10, 19), datetime(2026, 10, 19, 9, 0), listar=_listar(), contar_backlog=lambda f, a: 0)
    assert info["novos"] == 0 and info["registros"] == 0 and info["caminho"] is None
    assert len(os.listdir(tmp_path / "CACT - PUSH Diário")) == 1


def test_sem_email_elegivel_nao_gera_nem_grava_razao(cfg, tmp_path):
    _cfg(cfg, tmp_path)
    info = pc_.coletar(cfg, date(2026, 10, 19), AGORA, listar=lambda f, d: [EMAILS[2]], contar_backlog=lambda f, a: 0)
    assert info["elegiveis"] == 0 and info["caminho"] is None and not os.path.exists(cfg["push"]["ledger"])


def test_arquivar_so_quando_configurado_e_depois_do_razao(cfg, tmp_path):
    movidos = []

    def mover(emails, destino):
        movidos.append((sorted(e.message_id for e in emails), destino, os.path.exists(cfg["push"]["ledger"])))

    _cfg(cfg, tmp_path)
    pc_.coletar(cfg, date(2026, 10, 19), AGORA, listar=_listar(), contar_backlog=lambda f, a: 0, mover=mover)
    assert movidos == []                                                                        # arquivar_em = None
    _cfg(cfg, tmp_path / "b", arquivar_em="Tribunais/PUSH")
    os.makedirs(tmp_path / "b" / "CACT - PUSH Diário", exist_ok=True)
    pc_.coletar(cfg, date(2026, 10, 19), AGORA, listar=_listar(), contar_backlog=lambda f, a: 0, mover=mover)
    assert movidos == [(["<jus1@jusbrasil>", "<pje1@trt18>"], "Tribunais/PUSH", True)]         # só os que entraram no relatório; razão já gravado


def test_dry_run_nao_grava_razao_nem_move(cfg, tmp_path):
    _cfg(cfg, tmp_path, arquivar_em="Tribunais/PUSH")
    movidos = []
    info = pc_.coletar(cfg, date(2026, 10, 19), AGORA, listar=_listar(), contar_backlog=lambda f, a: 0,
                       mover=lambda e, d: movidos.append(d), dry_run=True, destino=str(tmp_path / "ensaio.xlsx"))
    assert info["registros"] == 5 and info["caminho"] == str(tmp_path / "ensaio.xlsx") and os.path.exists(info["caminho"])
    assert movidos == [] and not os.path.exists(cfg["push"]["ledger"]) and caminho_push(cfg, date(2026, 10, 19)) is None


def test_email_sem_evento_extraido_fica_de_fora_do_razao(cfg, tmp_path):
    _cfg(cfg, tmp_path)
    vazio = Email("sistema-x@trt18.jus.br", "[PUSH] aviso sem corpo útil", "", datetime(2026, 10, 16, 8, 0), "<vazio@trt18>", "x", "Caixa de Entrada")
    info = pc_.coletar(cfg, date(2026, 10, 19), AGORA, listar=lambda f, d: [EMAILS[0], vazio], contar_backlog=lambda f, a: 0)
    razao = json.load(open(cfg["push"]["ledger"], encoding="utf-8"))
    assert info["sem_registro"] == ["<vazio@trt18>"] and "<vazio@trt18>" not in razao and "<pje1@trt18>" in razao


# --------------------------------------------------------------------------- vigia
def test_passo_vigia_coleta_uma_vez_por_dia_util_antes_da_paralela(cfg, tmp_path, monkeypatch):
    from pc import agenda_paralela as ap
    _cfg(cfg, tmp_path)
    cfg["subpasta_paralela"] = "CACT - Agenda Paralela"
    os.makedirs(tmp_path / "CACT - Agenda Paralela", exist_ok=True)
    monkeypatch.setattr(ap, "avisar", lambda m: None)
    chamadas = []

    def coletar_falso(cfg_, dia, agora):
        chamadas.append((dia, agora))
        return {"novos": 0, "registros": 0, "caminho": None, "elegiveis": 0, "backlog": 0, "sem_registro": []}

    monkeypatch.setattr(ap, "coletar", coletar_falso)
    estado = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 19), datetime(2026, 10, 19, 8, 30), estado)      # Paralela ainda não chegou: coleta mesmo assim
    ap.passo_vigia(cfg, date(2026, 10, 19), datetime(2026, 10, 19, 8, 45), estado)
    assert chamadas == [(date(2026, 10, 19), datetime(2026, 10, 19, 8, 30))] and estado.coleta_feita is True
    estado2 = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 17), datetime(2026, 10, 17, 9, 0), estado2)        # sábado: não coleta
    assert len(chamadas) == 1


def test_passo_vigia_repete_coleta_se_outlook_falhar(cfg, tmp_path, monkeypatch):
    from pc import agenda_paralela as ap
    _cfg(cfg, tmp_path)
    monkeypatch.setattr(ap, "avisar", lambda m: None)
    tentativas = []

    def coletar_falho(cfg_, dia, agora):
        tentativas.append(agora)
        raise RuntimeError("conta não encontrada no Outlook")

    monkeypatch.setattr(ap, "coletar", coletar_falho)
    estado = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 19), datetime(2026, 10, 19, 8, 30), estado)
    ap.passo_vigia(cfg, date(2026, 10, 19), datetime(2026, 10, 19, 8, 45), estado)
    assert len(tentativas) == 2 and estado.coleta_feita is False


def test_passo_vigia_nao_coleta_se_relatorio_do_dia_ja_existe(cfg, tmp_path, monkeypatch):
    from pc import agenda_paralela as ap
    _cfg(cfg, tmp_path)
    monkeypatch.setattr(ap, "avisar", lambda m: None)
    open(os.path.join(tmp_path, "CACT - PUSH Diário", "Relatorio_PUSH_2026-10-19_0900.xlsx"), "wb").close()
    monkeypatch.setattr(ap, "coletar", lambda *a: (_ for _ in ()).throw(AssertionError("não devia coletar")))
    estado = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 19), datetime(2026, 10, 19, 9, 15), estado)
    assert estado.coleta_feita is True
