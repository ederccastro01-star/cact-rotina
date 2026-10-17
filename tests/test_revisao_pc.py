# -*- coding: utf-8 -*-
"""Testes da rodada de correções da revisão — lado do PC (C1, C2, C3, I1, I2, I3, I8, I9)."""
import os
from datetime import date, datetime

import openpyxl
import pytest
import pywintypes

from paralela.planilha import Linha, gerar_paralela
from pc import agenda_paralela as ap, excel_com, transporte as tr


def _cfg_pc(cfg, tmp_path):
    cfg["pc"] = {"pasta_controladoria": str(tmp_path), "hora_alerta_sem_arquivo": "10:00", "intervalo_min": 15}
    cfg["subpasta_paralela"] = "CACT - Agenda Paralela"
    os.makedirs(tmp_path / "CACT - Agenda Paralela", exist_ok=True)
    return cfg


def _linha(ids, prazo=date(2026, 10, 29)):
    return Linha(prazo=prazo, peca="Contrarrazões", publicado=date(2026, 10, 19), processo="0000001-11.2026.5.18.0001",
                 reclamante="A", reclamado="R", advogado="E", prioridade="NORMAL", cliente="AUTOR(A)", despacho="resumo",
                 destino="PRAZO DO ESCRITÓRIO", ids=ids, link="https://x/" + ids)


def _paralela(cfg, dia, linhas, decisoes=None, aba="Lançamentos"):
    caminho = ap.caminho_paralela(cfg, dia)
    gerar_paralela(linhas, [], caminho, dia, {"ok": True})
    if decisoes:
        wb = openpyxl.load_workbook(caminho); ws = wb[aba]
        for linha, dec in decisoes.items():
            ws.cell(linha, 13, dec)
        wb.save(caminho)
    return caminho


# ----------------------------------------------------------------------------- C1: Pendentes
def test_pendente_tem_origem_em_observacoes_e_lote_vazio(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    _paralela(cfg, date(2026, 10, 15), [_linha("1")])
    hoje = _paralela(cfg, date(2026, 10, 16), [_linha("3")])
    ap.montar_pendentes(cfg, date(2026, 10, 16))
    ws = openpyxl.load_workbook(hoje)["Pendentes"]
    assert ws["X8"].value == "Paralela 15-10-2026 (sem decisão)" and ws["W8"].value in (None, "")
    linhas = tr.ler_paralela(hoje)
    pend = [l for l in linhas if l["aba"] == "Pendentes"]
    assert pend and pend[0]["lote"] == "" and pend[0]["origem"] == "Paralela 15-10-2026 (sem decisão)"


def test_pendente_agendar_e_selecionada(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    _paralela(cfg, date(2026, 10, 15), [_linha("1")])
    hoje = _paralela(cfg, date(2026, 10, 16), [_linha("3")])
    ap.montar_pendentes(cfg, date(2026, 10, 16))
    wb = openpyxl.load_workbook(hoje); wb["Pendentes"]["M8"] = "Agendar"; wb.save(hoje)
    sel, avisos = tr.selecionar(tr.ler_paralela(hoje), set())
    assert [l["ids"] for l in sel] == [["1"]] and avisos == []


def test_decisao_em_pendentes_nao_ressuscita_amanha(cfg, tmp_path, monkeypatch):
    """D1 sem decisão; em D2 o advogado decide na aba Pendentes; em D3 a linha não volta."""
    _cfg_pc(cfg, tmp_path); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    _paralela(cfg, date(2026, 10, 14), [_linha("1"), _linha("2")])
    d2 = _paralela(cfg, date(2026, 10, 15), [_linha("3")])
    ap.montar_pendentes(cfg, date(2026, 10, 15))
    wb = openpyxl.load_workbook(d2); ws = wb["Pendentes"]
    alvo = next(r for r in range(8, 10) if ws.cell(r, 22).value == "1"); ws.cell(alvo, 13, "Descartar"); wb.save(d2)
    _paralela(cfg, date(2026, 10, 16), [_linha("4")])
    pend = ap.montar_pendentes(cfg, date(2026, 10, 16))
    assert sorted(l["ids"][0] for l in pend) == ["2", "3"]          # "1" foi decidida em Pendentes; não volta


def test_ids_na_agenda_consideram_prazos_feitos_e_razoes(cfg, tmp_path):
    _cfg_pc(cfg, tmp_path)
    r = tr.Razao(quando="x", paralela="p", linhas=[{"aba": "Publicações", "linha_agenda": 1, "ids": ["9"], "valores": {}, "linha_paralela": 8}])
    r.salvar(str(tmp_path / "CACT - Agenda Paralela" / "razao" / "a.json"))
    colunas = {"Publicações": ["1"], "Prazos Feitos": ["2;3"]}
    ids = ap.ids_ja_lancados(cfg, lambda aba: colunas[aba])
    assert ids == {"1", "2", "3", "9"}


# ----------------------------------------------------------------------------- C2: Excel ocupado no meio do lote
class _Cells:
    def __init__(self, ws, r, c): self.ws, self.r, self.c = ws, r, c
    @property
    def Value(self): return self.ws.dados.get((self.r, self.c))
    @Value.setter
    def Value(self, v): self.ws.dados[(self.r, self.c)] = v; self.ws.eventos.append(("set", self.r, self.c, v))


class _Ws:
    Name = "Publicações"
    def __init__(self, eventos): self.dados, self.eventos, self.proxima = {}, eventos, 3230
    def Cells(self, r, c): return _Cells(self, r, c)


class _Wb:
    def __init__(self, ws): self.ws = ws
    def Worksheets(self, nome): return self.ws
    def Save(self): self.ws.eventos.append(("Save",))


def _fakes(monkeypatch, eventos, ws, rejeitar_vezes=0):
    estado = {"n": rejeitar_vezes}

    def set_cell(ws_, r, c, v):
        if estado["n"] > 0:
            estado["n"] -= 1
            raise pywintypes.com_error(-2147418111, "ocupado", None, None)
        ws_.Cells(r, c).Value = v

    monkeypatch.setattr(excel_com.time, "sleep", lambda s: None)
    monkeypatch.setattr(ap, "abrir_livro", lambda c, visivel=False: (object(), _Wb(ws), False))
    monkeypatch.setattr(ap, "proxima_linha_tabela", lambda ws_, nome="Tabela1": (setattr(ws_, "proxima", ws_.proxima + 1), ws_.proxima)[1])
    monkeypatch.setattr(ap, "coluna_valores", lambda ws_, col, primeira=8: [])
    monkeypatch.setattr(ap, "set_cell", set_cell); monkeypatch.setattr(ap, "set_text", set_cell)
    monkeypatch.setattr(ap, "garantir_colunas_controle", lambda ws_, cfg_: None)
    monkeypatch.setattr(ap, "avisar", lambda m: None); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    monkeypatch.setattr(ap, "_salvar_razao", lambda razao, cfg_: eventos.append(("razao_salvo", razao.tipo, len(razao.linhas))) or "r.json")


def test_transportar_espera_excel_ocupado_e_termina_o_lote(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); eventos = []; ws = _Ws(eventos)
    _fakes(monkeypatch, eventos, ws, rejeitar_vezes=3)
    par = _paralela(cfg, date(2026, 10, 16), [_linha("1"), _linha("2")], decisoes={8: "Agendar", 9: "Agendar"})
    razao = ap.transportar(cfg, par, str(tmp_path / "agenda.xlsx"))
    assert razao and len(razao.linhas) == 2
    gravados = {(e[1], e[2]) for e in eventos if e[0] == "set"}
    assert (3231, 15) in gravados and (3232, 15) in gravados and ("Save",) in eventos


def test_razao_em_andamento_antes_da_primeira_escrita(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); eventos = []; ws = _Ws(eventos)
    _fakes(monkeypatch, eventos, ws)
    par = _paralela(cfg, date(2026, 10, 16), [_linha("1")], decisoes={8: "Agendar"})
    ap.transportar(cfg, par, str(tmp_path / "agenda.xlsx"))
    nomes = [e[0] for e in eventos]
    assert nomes.index("razao_salvo") < nomes.index("set")                      # razão "em andamento" antes de tocar na Agenda
    assert [e[1] for e in eventos if e[0] == "razao_salvo"] == ["em andamento", "transporte"]
    assert nomes.index("Save") > max(i for i, n in enumerate(nomes) if n == "razao_salvo")


# ----------------------------------------------------------------------------- C3: tabela até S
def test_garantir_colunas_controle_redimensiona_tabela_e_poe_cabecalhos():
    eventos = []

    class Rng:
        Row, Rows = 7, type("R", (), {"Count": 3221})()
        class Columns: Count = 11
        def Resize(self, r, c): eventos.append(("Resize", r, c)); return "novo"

    class Tab:
        Range = Rng()
        def Resize(self, rng): eventos.append(("ResizeTab", rng))

    class Ws:
        def __init__(self): self.dados = {(7, 14): "Polo conferido"}
        def ListObjects(self, nome): return Tab()
        def Cells(self, r, c):
            ws = self
            class C:
                @property
                def Value(s): return ws.dados.get((r, c))
                @Value.setter
                def Value(s, v): ws.dados[(r, c)] = v; eventos.append(("hdr", r, c, v))
                Font = type("F", (), {"Bold": False})()
            return C()
        def Range(self, a, b): return Rng()

    ws = Ws()
    ap.garantir_colunas_controle(ws, {"tabela_publicacoes": "Tabela1"}, vazias=lambda ws_, c1, c2: True)
    assert ("ResizeTab", "novo") in eventos and ("Resize", 3221, 19) in eventos
    assert ws.dados[(7, 15)] == "ID DJEN" and ws.dados[(7, 19)] == "CONFERIR" and ws.dados[(7, 14)] == "Polo conferido"


def test_garantir_colunas_controle_recusa_se_ha_dados_em_O_S():
    class Rng:
        Row, Rows = 7, type("R", (), {"Count": 3221})()
        class Columns: Count = 11
    class Tab:
        Range = Rng()
    class Ws:
        def ListObjects(self, nome): return Tab()
    with pytest.raises(RuntimeError):
        ap.garantir_colunas_controle(Ws(), {"tabela_publicacoes": "Tabela1"}, vazias=lambda ws_, c1, c2: False)


# ----------------------------------------------------------------------------- I1: um vigia só
def test_trava_impede_segundo_vigia(tmp_path):
    t1 = ap.obter_trava(str(tmp_path / "vigia.lock"))
    assert t1 is not None
    assert ap.obter_trava(str(tmp_path / "vigia.lock")) is None
    ap.soltar_trava(t1)
    assert ap.obter_trava(str(tmp_path / "vigia.lock")) is not None


# ----------------------------------------------------------------------------- I2: decisões antigas não somem
def test_paralelas_com_decisao_sem_lote(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    ontem = _paralela(cfg, date(2026, 10, 15), [_linha("1")], decisoes={8: "Agendar"})     # decidida, nunca transportada
    _paralela(cfg, date(2026, 10, 14), [_linha("2")])                                       # sem decisão
    assert ap.paralelas_com_decisao_sem_lote(cfg, date(2026, 10, 16)) == [ontem]


def test_passo_vigia_transporta_paralelas_anteriores(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    monkeypatch.setattr(ap, "avisar", lambda m: None); monkeypatch.setattr(ap, "abrir_para_o_advogado", lambda c, n=0: None)
    transportadas = []
    monkeypatch.setattr(ap, "transportar", lambda cfg_, par, agenda: transportadas.append(os.path.basename(par)))
    ontem = _paralela(cfg, date(2026, 10, 15), [_linha("1")], decisoes={8: "Agendar"})
    hoje = _paralela(cfg, date(2026, 10, 16), [_linha("3")])
    monkeypatch.setattr(ap, "arquivo_pronto", lambda c, agora=None: True)
    estado = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 16), datetime(2026, 10, 16, 9, 0), estado)
    assert os.path.basename(ontem) in transportadas


# ----------------------------------------------------------------------------- I8: arquivo só-online
def test_arquivo_pronto_hidrata_so_online(tmp_path, monkeypatch):
    f = tmp_path / "p.xlsx"; f.write_bytes(b"x" * 10)
    seq = [0x400000, 0x20]
    monkeypatch.setattr(ap, "_atributos", lambda c: seq.pop(0))
    hidratados = []
    monkeypatch.setattr(ap, "_hidratar", lambda c: hidratados.append(c))
    assert ap.arquivo_pronto(str(f), agora=datetime.now().timestamp() + 100) is True
    assert hidratados == [str(f)]


# ----------------------------------------------------------------------------- I9: data calculada preservada
def test_planilha_guarda_prazo_calculado_em_Y(tmp_path):
    caminho = str(tmp_path / "p.xlsx")
    gerar_paralela([_linha("1", prazo=date(2026, 10, 29))], [], caminho, date(2026, 10, 16), {"ok": True})
    ws = openpyxl.load_workbook(caminho)["Lançamentos"]
    assert ws["Y7"].value == "Prazo calculado" and ws["Y8"].value.date() == date(2026, 10, 29)


def test_comparar_usa_prazo_calculado_e_le_prazos_feitos(tmp_path):
    from pc import comparar_manual as cm
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Publicações"; ws.cell(7, 1, "Prazo")
    pf = wb.create_sheet("Prazos Feitos"); pf.cell(7, 1, "Prazo")
    pf.cell(8, 1, date(2026, 10, 26)); pf.cell(8, 3, "CMAP"); pf.cell(8, 4, date(2026, 10, 19)); pf.cell(8, 5, "0000001-11.2026.5.18.0001")
    agenda = str(tmp_path / "agenda.xlsx"); wb.save(agenda)
    par = str(tmp_path / "CACT - Agenda Paralela 16-10-2026.xlsx")
    gerar_paralela([_linha("1", prazo=date(2026, 10, 29))], [], par, date(2026, 10, 16), {"ok": True})
    w = openpyxl.load_workbook(par); w["Lançamentos"]["A8"] = date(2026, 10, 20); w["Lançamentos"]["M8"] = "Corrigir e agendar"; w.save(par)
    r = cm.comparar(agenda, [par], date(2026, 10, 15), date(2026, 10, 22))
    assert r["manuais"] == 1 and r["casadas"] == 1
    assert r["data_auto_posterior"] == [("0000001-11.2026.5.18.0001", date(2026, 10, 29), date(2026, 10, 26))]   # Y, não a A corrigida


# ----------------------------------------------------------------------------- I3: reprocessamento não sobrescreve
def test_nome_destino_paralela_com_sufixo_quando_ja_existe():
    import runner_nuvem as rn
    assert rn.nome_destino_paralela("CACT - Agenda Paralela 16-10-2026.xlsx", existe=False, agora=datetime(2026, 10, 16, 16, 5)) == "CACT - Agenda Paralela 16-10-2026.xlsx"
    assert rn.nome_destino_paralela("CACT - Agenda Paralela 16-10-2026.xlsx", existe=True, agora=datetime(2026, 10, 16, 16, 5)) == "CACT - Agenda Paralela 16-10-2026 (reprocessada 16h05).xlsx"
