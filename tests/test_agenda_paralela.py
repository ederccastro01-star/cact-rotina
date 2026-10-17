# -*- coding: utf-8 -*-
import os
from datetime import date, datetime

import openpyxl
import pytest

from paralela.planilha import Linha, gerar_paralela
from pc import agenda_paralela as ap


def _cfg_pc(cfg, tmp_path):
    cfg["pc"] = {"pasta_controladoria": str(tmp_path), "hora_alerta_sem_arquivo": "10:00", "intervalo_min": 15}
    cfg["subpasta_paralela"] = "CACT - Agenda Paralela"
    os.makedirs(tmp_path / "CACT - Agenda Paralela", exist_ok=True)
    return cfg


def _linha(ids, decisao_col=None):
    return Linha(prazo=date(2026, 10, 29), peca="Contrarrazões", publicado=date(2026, 10, 19), processo="0000001-11.2026.5.18.0001",
                 reclamante="A", reclamado="R", advogado="E", prioridade="NORMAL", cliente="AUTOR(A)", despacho="resumo",
                 destino="PRAZO DO ESCRITÓRIO", ids=ids, link="https://x/" + ids)


def _paralela(cfg, dia, linhas, decisoes=None):
    caminho = ap.caminho_paralela(cfg, dia)
    gerar_paralela(linhas, [], caminho, dia, {"ok": True})
    if decisoes:
        wb = openpyxl.load_workbook(caminho); ws = wb["Lançamentos"]
        for linha, dec in decisoes.items():
            ws.cell(linha, 13, dec)
        wb.save(caminho)
    return caminho


def test_caminho_paralela(cfg, tmp_path):
    _cfg_pc(cfg, tmp_path)
    assert ap.caminho_paralela(cfg, date(2026, 10, 16)) == str(tmp_path / "CACT - Agenda Paralela" / "CACT - Agenda Paralela 16-10-2026.xlsx")


def test_alerta_sem_arquivo_as_10h(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path)
    avisos = []
    monkeypatch.setattr(ap, "avisar", lambda m: avisos.append(m))
    estado = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 16), datetime(2026, 10, 16, 9, 30), estado)
    assert avisos == []
    ap.passo_vigia(cfg, date(2026, 10, 16), datetime(2026, 10, 16, 10, 1), estado)
    ap.passo_vigia(cfg, date(2026, 10, 16), datetime(2026, 10, 16, 10, 16), estado)
    assert len(avisos) == 1 and "VM não entregou" in avisos[0]


def test_arquivo_online_only_nao_esta_pronto(cfg, tmp_path, monkeypatch):
    f = tmp_path / "p.xlsx"; f.write_bytes(b"x" * 10)
    monkeypatch.setattr(ap, "_atributos", lambda c: 0x400000)
    assert ap.arquivo_pronto(str(f), agora=datetime.now().timestamp() + 100) is False
    monkeypatch.setattr(ap, "_atributos", lambda c: 0x20)
    assert ap.arquivo_pronto(str(f), agora=datetime.now().timestamp() + 100) is True
    assert ap.arquivo_pronto(str(f), agora=datetime.now().timestamp()) is False      # mtime ainda não estável (10 s)
    assert ap.arquivo_pronto(str(tmp_path / "nao_existe.xlsx")) is False


def test_montar_pendentes_ignora_decididas(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path)
    monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    _paralela(cfg, date(2026, 10, 15), [_linha("1"), _linha("2")], decisoes={8: "Agendar"})
    hoje = _paralela(cfg, date(2026, 10, 16), [_linha("3")])
    pend = ap.montar_pendentes(cfg, date(2026, 10, 16))
    assert len(pend) == 1 and pend[0]["ids"] == ["2"]
    ws = openpyxl.load_workbook(hoje)["Pendentes"]
    assert ws["E8"].value == "0000001-11.2026.5.18.0001" and ws["V8"].value == "2"
    assert ws["W8"].value == "Paralela 15-10-2026 (sem decisão)"


class _Cells:
    def __init__(self, ws, r, c):
        self.ws, self.r, self.c = ws, r, c

    @property
    def Value(self):
        return self.ws.dados.get((self.r, self.c))

    @Value.setter
    def Value(self, v):
        self.ws.dados[(self.r, self.c)] = v; self.ws.eventos.append(("set", self.r, self.c, v))


class _Ws:
    def __init__(self, eventos, ids_col_o=()):
        self.dados, self.eventos, self.proxima = {}, eventos, 3230
        for i, v in enumerate(ids_col_o):
            self.dados[(8 + i, 15)] = v
        self.Rows = type("R", (), {"Count": 1048576})()

    def Cells(self, r, c):
        return _Cells(self, r, c)


class _Wb:
    def __init__(self, ws):
        self.ws, self.Saved = ws, True

    def Worksheets(self, nome):
        return self.ws

    def Save(self):
        self.ws.eventos.append(("Save",))


def _fake_excel(ids_col_o=()):
    eventos = []
    ws = _Ws(eventos, ids_col_o)

    def abrir(caminho, visivel=False):
        return object(), _Wb(ws), False

    def proxima(ws_, nome="Tabela1"):
        ws_.proxima += 1; return ws_.proxima

    def coluna(ws_, col, primeira=8):
        return [ws_.dados.get((r, col)) for r in range(8, 8 + len(ids_col_o))]

    return eventos, abrir, proxima, coluna


def test_transportar_grava_razao_antes_do_save(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path)
    eventos, abrir, proxima, coluna = _fake_excel()
    monkeypatch.setattr(ap, "abrir_livro", abrir); monkeypatch.setattr(ap, "proxima_linha_tabela", proxima)
    monkeypatch.setattr(ap, "coluna_valores", coluna); monkeypatch.setattr(ap, "set_cell", lambda ws, r, c, v: setattr(ws.Cells(r, c), "Value", v))
    monkeypatch.setattr(ap, "set_text", lambda ws, r, c, v: eventos.append(("set_text", r, c, v)))
    monkeypatch.setattr(ap, "avisar", lambda m: None); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    monkeypatch.setattr(ap, "_salvar_razao", lambda razao, cfg_: eventos.append(("razao_salvo", len(razao.linhas))) or "r.json")
    paralela = _paralela(cfg, date(2026, 10, 16), [_linha("1"), _linha("2")], decisoes={8: "Agendar", 9: "Descartar"})
    razao = ap.transportar(cfg, paralela, str(tmp_path / "agenda.xlsx"))
    assert razao is not None and len(razao.linhas) == 1
    nomes = [e[0] for e in eventos]
    assert nomes.index("razao_salvo") < nomes.index("Save")
    gravados = {(e[1], e[2]): e[3] for e in eventos if e[0] == "set"}
    textos = {(e[1], e[2]): e[3] for e in eventos if e[0] == "set_text"}
    assert textos[(3231, 15)] == "1" and textos[(3231, 5)] == "0000001-11.2026.5.18.0001"   # ID e processo como TEXTO
    assert gravados[(3231, 17)] == "AUTO"
    assert gravados[(3231, 2)] == '=IF(A3231="","",A3231)'
    assert (3231, 14) not in gravados
    ws = openpyxl.load_workbook(paralela)["Lançamentos"]
    assert str(ws["W8"].value).startswith("Transportado em") and ws["W9"].value in (None, "")


def test_transportar_segunda_vez_zero_linhas(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path)
    eventos, abrir, proxima, coluna = _fake_excel(ids_col_o=["1"])
    monkeypatch.setattr(ap, "abrir_livro", abrir); monkeypatch.setattr(ap, "proxima_linha_tabela", proxima)
    monkeypatch.setattr(ap, "coluna_valores", coluna); monkeypatch.setattr(ap, "avisar", lambda m: None)
    monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    paralela = _paralela(cfg, date(2026, 10, 16), [_linha("1")], decisoes={8: "Agendar"})
    assert ap.transportar(cfg, paralela, str(tmp_path / "agenda.xlsx")) is None
    assert ("Save",) not in eventos


def test_desfazer_apaga_linha_inteira(cfg, tmp_path, monkeypatch):
    """ListRows.Delete só apaga A–K (a Tabela1 vai até K) e deixa O–S com os IDs: tem de apagar a linha da planilha."""
    _cfg_pc(cfg, tmp_path)
    eventos = []

    class Linha_:
        def __init__(self, r): self.r = r
        def Delete(self): eventos.append(("Rows.Delete", self.r))

    class Ws:
        Name = "Publicações"
        def Rows(self, r): return Linha_(r)
        def ListObjects(self, nome): raise AssertionError("não usar ListRows.Delete no desfazer")

    class Wb:
        def Worksheets(self, nome): return Ws()
        def Save(self): eventos.append(("Save",))

    monkeypatch.setattr(ap, "abrir_livro", lambda c, visivel=False: (object(), Wb(), False))
    monkeypatch.setattr(ap, "coluna_valores", lambda ws, col, primeira=8: [None, "111", None, "222"])   # linhas 9 e 11
    monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    razao = ap.Razao(quando="x", paralela=str(tmp_path / "inexistente.xlsx"),
                     linhas=[{"aba": "Publicações", "linha_agenda": 3228, "ids": ["111"], "valores": {}, "linha_paralela": 8},
                             {"aba": "Publicações", "linha_agenda": 3229, "ids": ["222"], "valores": {}, "linha_paralela": 9}])
    rz = tmp_path / "r.json"; razao.salvar(str(rz))
    assert ap.desfazer(cfg, str(rz)) == 2
    assert eventos == [("Rows.Delete", 11), ("Rows.Delete", 9), ("Save",)]      # de baixo para cima, depois salva
