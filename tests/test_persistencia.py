# -*- coding: utf-8 -*-
"""Lote transportado que sumiu inteiro da Agenda (conflito de versão do OneDrive, por exemplo) tem de voltar a transportar
sozinho. A busca é pelo ID em todas as abas: linha que só mudou de posição (tabela ordenada) não é perda."""
import os
from datetime import date, datetime

import openpyxl

from paralela.planilha import Linha, gerar_paralela
from pc import agenda_paralela as ap
from pc.transporte import COL_PARALELA, Razao


def _cfg(cfg, tmp_path):
    cfg["pc"] = {"pasta_controladoria": str(tmp_path), "hora_alerta_sem_arquivo": "10:00", "intervalo_min": 15}
    cfg["subpasta_paralela"] = "CACT - Agenda Paralela"
    os.makedirs(tmp_path / "CACT - Agenda Paralela" / "razao", exist_ok=True)
    return cfg


def _linha(ids, processo):
    return Linha(prazo=date(2026, 10, 21), peca="Manifestação", publicado=date(2026, 10, 16), processo=processo, reclamante="A",
                 reclamado="R", advogado="E", prioridade="NORMAL", cliente="AUTOR(A)", despacho="r", destino="PRAZO DO ESCRITÓRIO", ids=ids)


def _paralela_transportada(cfg, dia, ids, lote="Transportado em 17/10/2026 19:06"):
    par = ap.caminho_paralela(cfg, dia)
    gerar_paralela([_linha(";".join(ids), "0010097-61.2018.5.18.0015")], [], par, dia, {"ok": True})
    wb = openpyxl.load_workbook(par); ws = wb["Lançamentos"]
    ws.cell(8, COL_PARALELA["decisao"], "Agendar"); ws.cell(8, COL_PARALELA["lote"], lote); wb.save(par)
    return par


def _razao(cfg, par, ids, quando="2026-10-17T19:06:37", tipo="transporte"):
    r = Razao(quando=quando, paralela=par, tipo=tipo,
              linhas=[{"aba": "Publicações", "linha_agenda": 3228, "ids": ids, "valores": {}, "linha_paralela": 8, "aba_paralela": "Lançamentos"}])
    c = os.path.join(ap.pasta_paralela(cfg), "razao", quando.replace("-", "").replace(":", "")[:15] + ".json")
    r.salvar(c)
    return c


def _agenda(cfg, ids_publicacoes=(), ids_feitos=(), ids_arquivo=()):
    caminho = ap.caminho_agenda(cfg)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Publicações"
    for i, v in enumerate(ids_publicacoes, 8):
        ws.cell(i, 5, "0000001-11.2026.5.18.0001"); ws.cell(i, 15, v)
    pf = wb.create_sheet("Prazos Feitos")
    for i, v in enumerate(ids_feitos, 8):
        pf.cell(i, 5, "0000002-11.2026.5.18.0002"); pf.cell(i, 15, v)
    arq = wb.create_sheet("Prazos Feitos (Arquivo)")
    for i, v in enumerate(ids_arquivo, 8):
        arq.cell(i, 5, "0000003-11.2026.5.18.0003"); arq.cell(i, 15, v)
    wb.save(caminho)
    return caminho


def test_ids_agenda_por_copia_le_todas_as_abas(cfg, tmp_path):
    _cfg(cfg, tmp_path); _agenda(cfg, ["111;222"], ["333"], ["444"])
    assert ap.ids_agenda_por_copia(cfg) == {"111", "222", "333", "444"}


def test_lote_sumido_da_agenda_volta_a_transportar(cfg, tmp_path, monkeypatch):
    _cfg(cfg, tmp_path); monkeypatch.setattr(ap, "avisar", lambda m: None); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    par = _paralela_transportada(cfg, date(2026, 10, 15), ["756801624", "756801644"])
    c = _razao(cfg, par, ["756801624", "756801644"])
    perdidos = ap.verificar_persistencia(cfg, date(2026, 10, 18), ids_fn=lambda: set())
    assert perdidos == [c] and Razao.carregar(c).tipo == "perdido"
    ws = openpyxl.load_workbook(par)["Lançamentos"]
    assert not ws.cell(8, COL_PARALELA["lote"]).value and "perdido" in str(ws.cell(8, COL_PARALELA["obs"]).value)
    assert not ({"756801624", "756801644"} & ap.ids_ja_lancados(cfg, lambda aba: []))       # razão perdido não bloqueia mais
    assert ap.paralelas_com_decisao_sem_lote(cfg, date(2026, 10, 18)) == [par]                 # vigia vai transportar de novo
    assert ap.verificar_persistencia(cfg, date(2026, 10, 18), ids_fn=lambda: set()) == []     # não repete


def test_lote_presente_ou_parcial_nao_e_tocado(cfg, tmp_path, monkeypatch):
    _cfg(cfg, tmp_path); monkeypatch.setattr(ap, "avisar", lambda m: None); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    par = _paralela_transportada(cfg, date(2026, 10, 15), ["1", "2"])
    c = _razao(cfg, par, ["1", "2"])
    assert ap.verificar_persistencia(cfg, date(2026, 10, 18), ids_fn=lambda: {"2"}) == []      # advogado pode ter apagado uma só
    assert Razao.carregar(c).tipo == "transporte" and openpyxl.load_workbook(par)["Lançamentos"].cell(8, COL_PARALELA["lote"]).value


def test_lote_antigo_e_tipos_nao_transporte_sao_ignorados(cfg, tmp_path, monkeypatch):
    _cfg(cfg, tmp_path); monkeypatch.setattr(ap, "avisar", lambda m: None); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    par = _paralela_transportada(cfg, date(2026, 10, 1), ["9"])
    _razao(cfg, par, ["9"], quando="2026-10-01T10:00:00")                     # mais de 7 dias
    _razao(cfg, par, ["8"], quando="2026-10-17T10:00:00", tipo="desfazer")
    assert ap.verificar_persistencia(cfg, date(2026, 10, 18), ids_fn=lambda: set()) == []


def test_passo_vigia_verifica_persistencia_a_cada_15_min(cfg, tmp_path, monkeypatch):
    _cfg(cfg, tmp_path); monkeypatch.setattr(ap, "avisar", lambda m: None)
    chamadas = []
    monkeypatch.setattr(ap, "verificar_persistencia", lambda cfg_, dia, **k: chamadas.append(dia) or [])
    estado = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 18), datetime(2026, 10, 18, 10, 0), estado)
    ap.passo_vigia(cfg, date(2026, 10, 18), datetime(2026, 10, 18, 10, 5), estado)
    ap.passo_vigia(cfg, date(2026, 10, 18), datetime(2026, 10, 18, 10, 16), estado)
    assert chamadas == [date(2026, 10, 18), date(2026, 10, 18)]


def test_linha_que_o_advogado_passou_a_descartar_nao_e_perda(cfg, tmp_path, monkeypatch):
    """18/10/2026 12:01: linha transportada no automático; o advogado mudou a decisão para Descartar e apagou a linha da
    Agenda. Lote sem nenhuma linha ainda marcada para agendar não é perda: nada a retransportar, nenhum aviso."""
    _cfg(cfg, tmp_path); avisos = []; monkeypatch.setattr(ap, "avisar", lambda m: avisos.append(m)); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    par = _paralela_transportada(cfg, date(2026, 10, 16), ["758923180"], lote="Transportado em 18/10/2026 12:01")
    wb = openpyxl.load_workbook(par); wb["Lançamentos"].cell(8, COL_PARALELA["decisao"], "Descartar"); wb.save(par)
    c = _razao(cfg, par, ["758923180"], quando="2026-10-18T12:01:49")
    assert ap.verificar_persistencia(cfg, date(2026, 10, 18), ids_fn=lambda: set()) == []
    assert Razao.carregar(c).tipo == "transporte" and avisos == []
    assert openpyxl.load_workbook(par)["Lançamentos"].cell(8, COL_PARALELA["lote"]).value == "Transportado em 18/10/2026 12:01"
