# -*- coding: utf-8 -*-
from datetime import date, datetime

import openpyxl

from paralela.planilha import Linha, gerar_paralela
from pc import comparar_manual as cm


def _agenda(tmp_path):
    """Agenda falsa: cabeçalho na linha 7, dados da 8; 3 linhas manuais (Q vazia) e 1 automática (Q = AUTO)."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Publicações"
    ws.cell(7, 1, "Prazo"); ws.cell(7, 5, "Processo ")
    dados = [
        (date(2026, 10, 29), "CMAP", date(2026, 10, 19), "0000001-11.2026.5.18.0001", None),
        (date(2026, 10, 26), "IE", date(2026, 10, 19), "0000002-11.2026.5.18.0002", None),
        (date(2026, 10, 30), "RO", date(2026, 10, 20), "0000003-11.2026.5.18.0003", None),
        (date(2026, 10, 26), "ED", date(2026, 10, 19), "0000009-11.2026.5.18.0009", "AUTO"),
    ]
    for i, (prazo, peca, pub, proc, origem) in enumerate(dados, 8):
        ws.cell(i, 1, prazo); ws.cell(i, 3, peca); ws.cell(i, 4, pub); ws.cell(i, 5, proc)
        if origem:
            ws.cell(i, 17, origem)
    c = str(tmp_path / "agenda.xlsx"); wb.save(c); return c


def _paralela(tmp_path):
    def l(proc, prazo, peca="Contrarrazões", destino="PRAZO DO ESCRITÓRIO"):
        return Linha(prazo=prazo, peca=peca, publicado=date(2026, 10, 19), processo=proc, reclamante="A", reclamado="R",
                     advogado="E", prioridade="NORMAL", cliente="AUTOR(A)", despacho="r", destino=destino, ids=proc[:7])
    linhas = [l("0000001-11.2026.5.18.0001", date(2026, 10, 29)),                 # casa, mesma data
              l("0000002-11.2026.5.18.0002", date(2026, 10, 28)),                 # casa, data automática POSTERIOR (26/10 manual)
              l("0000003-11.2026.5.18.0003", date(2026, 10, 30)),                 # casa (publicado 20/10 dentro de ±4 dias)
              l("0000004-11.2026.5.18.0004", date(2026, 10, 26), destino="PARTE CONTRÁRIA")]  # só na Paralela
    c = str(tmp_path / "CACT - Agenda Paralela 16-10-2026.xlsx")
    gerar_paralela(linhas, [], c, date(2026, 10, 16), {"ok": True}); return c


def test_comparar(tmp_path):
    r = cm.comparar(_agenda(tmp_path), [_paralela(tmp_path)], date(2026, 10, 15), date(2026, 10, 22))
    assert r["manuais"] == 3                 # a linha AUTO não é manual
    assert r["casadas"] == 3
    assert r["data_auto_posterior"] == [("0000002-11.2026.5.18.0002", date(2026, 10, 28), date(2026, 10, 26))]
    assert r["so_na_paralela"] == 1 and r["so_na_agenda"] == 0


def test_relatorio_texto(tmp_path):
    r = cm.comparar(_agenda(tmp_path), [_paralela(tmp_path)], date(2026, 10, 15), date(2026, 10, 22))
    txt = cm.relatorio(r, date(2026, 10, 15), date(2026, 10, 22))
    assert "DATA AUTOMÁTICA POSTERIOR" in txt and "0000002-11.2026.5.18.0002" in txt and "manuais: 3" in txt
