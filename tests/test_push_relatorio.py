# -*- coding: utf-8 -*-
"""Relatorio_PUSH_*.xlsx gerado no PC: mesmo formato da rotina de nuvem (18 colunas + Resumo com fórmulas)."""
from datetime import datetime

import openpyxl

from pc import push_relatorio as pr
from pc.push_fontes import Registro

CAB = ["Data do andamento", "Data de autuação", "Nº do processo (CNJ)", "Tribunal", "Ramo da Justiça", "Grau", "Classe judicial",
       "Órgão / Vara", "Tipo de andamento", "Andamento / Teor", "Autor", "Réu", "Origem (sistema)", "Data do e-mail",
       "Assunto do e-mail", "Remetente", "ID da conversa", "Registro repetido?"]


def _reg(processo, data, teor, tipo="Intimação", origem="PJe Push", tribunal="TRT18", mid="<m1@x>"):
    return Registro(data=data, autuacao=datetime(2026, 5, 17, 11, 17, 44), processo=processo, tribunal=tribunal, ramo="Justiça do Trabalho",
                    grau="1º Grau", classe="ATOrd", orgao="11ª VT", tipo=tipo, teor=teor, autor="A", reu="R", origem=origem,
                    data_email=datetime(2026, 10, 16, 8, 52), assunto="ENC: [PUSH] x", remetente="nao-responda@trt18.jus.br", message_id=mid)


def _regs():
    return [
        _reg("0000001-11.2026.5.18.0011", datetime(2026, 10, 15, 23, 40), "Expedido(a) intimação a(o) A"),
        _reg("0000001-11.2026.5.18.0011", datetime(2026, 10, 16, 9, 0), "Documento sigiloso", tipo="Documento sigiloso", mid="<m2@x>"),
        _reg("0000001-11.2026.5.18.0011", datetime(2026, 10, 15, 23, 40), "Expedido(a) intimação a(o) A", mid="<m2@x>"),   # repetido (e-mail de outro dia)
        _reg("0000002-22.2024.5.18.0008", datetime(2026, 10, 16), "Pauta", tipo="Pauta/Audiência", origem="Jusbrasil", tribunal="TJGO", mid="<m3@x>"),
    ]


def _gerar(tmp_path, regs=None, backlog=0):
    caminho = str(tmp_path / "Relatorio_PUSH_2026-10-19_0845.xlsx")
    info = pr.gerar_relatorio(regs if regs is not None else _regs(), caminho, janela=(datetime(2026, 10, 12, 8, 45), datetime(2026, 10, 19, 8, 45)),
                              contas=["x@gmail.com", "y@hotmail.com"], backlog=backlog, gerado_em=datetime(2026, 10, 19, 8, 45))
    return caminho, info


def test_aba_principal_cabecalho_ordem_e_repetidos(tmp_path):
    caminho, info = _gerar(tmp_path)
    wb = openpyxl.load_workbook(caminho)
    assert wb.sheetnames == ["Publicações e Andamentos", "Resumo"]
    ws = wb["Publicações e Andamentos"]
    assert [c.value for c in ws[1]] == CAB
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 4 and info["registros"] == 4 and info["conversas"] == 3
    assert [r[0] for r in linhas] == sorted([r[0] for r in linhas], reverse=True)          # data desc
    rep = {(r[2], r[0], r[9]): [] for r in linhas}
    for r in linhas:
        rep[(r[2], r[0], r[9])].append(r[17])
    assert sorted(rep[("0000001-11.2026.5.18.0011", datetime(2026, 10, 15, 23, 40), "Expedido(a) intimação a(o) A")]) == ["Não", "Sim (2ª ocorrência)"]
    assert all(r[17] == "Não" for r in linhas if r[9] != "Expedido(a) intimação a(o) A")
    assert all(len(str(r[16])) == 16 for r in linhas)                                        # ID da conversa: hash curto do Message-ID


def test_formatacao_igual_a_da_rotina(tmp_path):
    caminho, _ = _gerar(tmp_path)
    ws = openpyxl.load_workbook(caminho)["Publicações e Andamentos"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:R5"
    assert ws["A1"].font.bold and ws["A1"].font.name == "Arial" and ws["A1"].fill.fgColor.rgb.endswith("1F3864")
    assert ws["A1"].font.color.rgb.endswith("FFFFFF") and ws.row_dimensions[1].height == 32
    assert ws["A2"].number_format == "DD/MM/YYYY HH:MM" and ws["N2"].number_format == "DD/MM/YYYY HH:MM"
    assert ws["J2"].alignment.wrap_text and ws["H2"].alignment.wrap_text
    assert ws["A3"].fill.fgColor.rgb.endswith("F2F2F2") and not (ws["A2"].fill.fgColor.rgb or "").endswith("F2F2F2")
    assert round(ws.column_dimensions["J"].width) == 90 and round(ws.column_dimensions["C"].width) == 26


def test_resumo_com_formulas_e_backlog(tmp_path):
    caminho, _ = _gerar(tmp_path, backlog=3)
    rs = openpyxl.load_workbook(caminho)["Resumo"]
    celulas = {str(c.value): c for linha in rs.iter_rows() for c in linha if c.value is not None}
    assert "Relatório de Publicações e Andamentos Processuais" in celulas
    assert any(v.startswith("Fonte: e-mails lidos no Outlook") and "x@gmail.com" in v for v in celulas)
    assert any(v.startswith("Janela coberta: 12/10/2026 08:45 a 19/10/2026 08:45") for v in celulas)
    total = [c for v, c in celulas.items() if v == "Total de registros"][0]
    assert rs.cell(total.row, 2).value == "=COUNTA('Publicações e Andamentos'!C2:C5)"
    conv = [c for v, c in celulas.items() if v == "Total de conversas processadas"][0]
    assert rs.cell(conv.row, 2).value == 3
    bl = [c for v, c in celulas.items() if v.startswith("Backlog")][0]
    assert rs.cell(bl.row, 2).value == 3
    assert "NOTAS DE QUALIDADE" in celulas and any(v.startswith("5. Backlog") and "3" in v for v in celulas)
    for bloco, col in (("POR TRIBUNAL", "D"), ("POR RAMO DA JUSTIÇA", "E"), ("POR TIPO DE ANDAMENTO", "I"), ("POR ORIGEM (SISTEMA)", "M")):
        cab = celulas[bloco]
        f = rs.cell(cab.row + 1, 2).value
        assert f.startswith("=COUNTIF('Publicações e Andamentos'!$" + col + "$2:$" + col + "$5,A") and rs.cell(cab.row + 1, 3).value.startswith("=B")
        assert rs.cell(cab.row + 1, 3).value.endswith("/$B$" + str(total.row))
    trib = celulas["POR TRIBUNAL"]
    assert rs.cell(trib.row + 1, 1).value == "TRT18" and rs.cell(trib.row + 2, 1).value == "TJGO"
    assert rs.cell(trib.row + 3, 1).value == "Total" and rs.cell(trib.row + 3, 2).value == f"=SUM(B{trib.row + 1}:B{trib.row + 2})"
    top = celulas["TOP 20 PROCESSOS COM MAIS REGISTROS"]
    assert rs.cell(top.row + 1, 1).value == "0000001-11.2026.5.18.0011" and rs.cell(top.row + 1, 2).value.startswith("=COUNTIF(")
    assert rs["A1"].font.name == "Arial"


def test_sem_registros_nao_gera_arquivo(tmp_path):
    caminho, info = _gerar(tmp_path, regs=[])
    import os
    assert not os.path.exists(caminho) and info["registros"] == 0
