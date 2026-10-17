# -*- coding: utf-8 -*-
from datetime import date

import openpyxl

from paralela.planilha import COLUNAS_AGENDA, COLUNAS_CONTROLE, Linha, gerar_paralela


def linha(destino="PRAZO DO ESCRITÓRIO", motivo="", prazo=date(2026, 10, 29), ids="1;2"):
    return Linha(prazo=prazo, peca="Contrarrazões", publicado=date(2026, 10, 19), processo="0000001-11.2026.5.18.0001",
                 reclamante="AUTOR TESTE", reclamado="REU TESTE LTDA", advogado="EDER", prioridade="NORMAL", cliente="AUTOR(A)",
                 despacho="Intimado: AUTOR TESTE (Autor) · Ato: decisão · Determinação: contrarrazões · Prazo: 8 dias (texto)",
                 destino=destino, motivo=motivo, polo_intimado="AUTOR", intimados="AUTOR TESTE", fonte_polo="DJEN + coluna J",
                 prazo_texto=8, prazo_legal=8, base_legal="CLT 895/897/896", link="https://x/1", ids=ids)


def test_gerar_paralela_layout(tmp_path):
    caminho = str(tmp_path / "CACT - Agenda Paralela 16-10-2026.xlsx")
    gerar_paralela([linha(), linha(destino="PARTE CONTRÁRIA", ids="3")], [linha(destino="FORA DA AGENDA", ids="4")],
                   caminho, date(2026, 10, 16), {"atos": 3, "processos": 1, "ids": 4, "ok": True})
    wb = openpyxl.load_workbook(caminho)
    assert wb.sheetnames == ["Leia-me", "Lançamentos", "Fora da Agenda", "Pendentes"]
    ws = wb["Lançamentos"]
    assert [ws.cell(7, i + 1).value for i in range(11)] == COLUNAS_AGENDA
    assert [ws.cell(7, 12 + i).value for i in range(len(COLUNAS_CONTROLE))] == COLUNAS_CONTROLE
    assert ws["L7"].value == "Destino" and ws["M7"].value == "Decisão" and ws["V7"].value == "ID DJEN"
    assert ws["B8"].value == '=IF(A8="","",A8)' and ws["B8"].number_format == "ddd"
    assert ws["E8"].font.name == "Courier New"
    assert ws["A8"].number_format == "dd/mm/yy;@"
    dv = ws.data_validations.dataValidation
    assert len(dv) == 1 and dv[0].formula1 == '"Agendar,Descartar,Corrigir e agendar"'
    assert "M8:M9" in str(dv[0].sqref)
    assert ws["M8"].fill.fgColor.rgb.endswith("FFFF00")
    assert ws["V8"].value == "1;2"
    assert ws.freeze_panes == "F8"
    assert ws["L8"].value == "PRAZO DO ESCRITÓRIO" and ws["L9"].value == "PARTE CONTRÁRIA"
    assert wb["Fora da Agenda"]["L8"].value == "FORA DA AGENDA"
    assert wb["Pendentes"]["A7"].value == "Prazo" and wb["Pendentes"]["A8"].value is None


def test_ordem_novo_conferir_escritorio_contraria(tmp_path):
    caminho = str(tmp_path / "p.xlsx")
    linhas = [linha(destino="PARTE CONTRÁRIA", ids="a"), linha(destino="PRAZO DO ESCRITÓRIO", prazo=date(2026, 10, 29), ids="b"),
              linha(destino="PRAZO DO ESCRITÓRIO", motivo="sem prazo", ids="c"), linha(destino="PROCESSO NOVO", motivo="novo", ids="d"),
              linha(destino="PRAZO DO ESCRITÓRIO", prazo=date(2026, 10, 26), ids="e")]
    gerar_paralela(linhas, [], caminho, date(2026, 10, 16), {"atos": 5, "processos": 1, "ids": 5, "ok": True})
    ws = openpyxl.load_workbook(caminho)["Lançamentos"]
    assert [ws.cell(r, 22).value for r in range(8, 13)] == ["d", "c", "e", "b", "a"]


def test_conciliacao_falha_vai_para_o_topo(tmp_path):
    caminho = str(tmp_path / "p.xlsx")
    gerar_paralela([linha()], [], caminho, date(2026, 10, 16), {"atos": 1, "processos": 1, "ids": 2, "ok": False, "faltam": ["99"]})
    ws = openpyxl.load_workbook(caminho)["Lançamentos"]
    assert "ATENÇÃO" in str(ws["A6"].value) and "99" in str(ws["A6"].value)
