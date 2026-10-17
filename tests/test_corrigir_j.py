# -*- coding: utf-8 -*-
"""Coluna Z 'Corrigir J para' na Paralela: o transporte corrige J em todas as linhas do processo na Agenda."""
import os
from datetime import date

import openpyxl

from paralela.planilha import COLUNAS_CONTROLE, Linha, gerar_paralela
from pc import agenda_paralela as ap, excel_com, transporte as tr


def _cfg_pc(cfg, tmp_path):
    cfg["pc"] = {"pasta_controladoria": str(tmp_path), "hora_alerta_sem_arquivo": "10:00", "intervalo_min": 15}
    cfg["subpasta_paralela"] = "CACT - Agenda Paralela"
    os.makedirs(tmp_path / "CACT - Agenda Paralela", exist_ok=True)
    return cfg


def _linha(ids, processo="0000001-11.2026.5.18.0001", cliente="AUTOR(A)", destino="PRAZO DO ESCRITÓRIO"):
    return Linha(prazo=date(2026, 10, 29), peca="Contrarrazões", publicado=date(2026, 10, 19), processo=processo,
                 reclamante="A", reclamado="R", advogado="E", prioridade="NORMAL", cliente=cliente, despacho="resumo",
                 destino=destino, ids=ids, link="https://x/" + ids)


def _paralela(cfg, dia, linhas, celulas=None):
    caminho = ap.caminho_paralela(cfg, dia)
    gerar_paralela(linhas, [], caminho, dia, {"ok": True})
    if celulas:
        wb = openpyxl.load_workbook(caminho); ws = wb["Lançamentos"]
        for ref, v in celulas.items():
            ws[ref] = v
        wb.save(caminho)
    return caminho


def test_planilha_tem_coluna_corrigir_j(tmp_path):
    caminho = str(tmp_path / "p.xlsx")
    gerar_paralela([_linha("1")], [], caminho, date(2026, 10, 16), {"ok": True})
    ws = openpyxl.load_workbook(caminho)["Lançamentos"]
    assert COLUNAS_CONTROLE[-1] == "Corrigir J para" and ws["Z7"].value == "Corrigir J para"
    assert ws["Z8"].fill.fgColor.rgb.endswith("FFFF00")
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert '"AUTOR(A),RÉU"' in formulas


def test_ler_paralela_le_corrigir_j_normalizado(cfg, tmp_path):
    _cfg_pc(cfg, tmp_path)
    par = _paralela(cfg, date(2026, 10, 16), [_linha("1"), _linha("2"), _linha("3")], {"Z8": "réu", "Z9": "autor", "Z10": "✔ RÉU (17/10 09:00)"})
    l = {x["ids"][0]: x for x in tr.ler_paralela(par)}
    assert l["1"]["corrigir_j"] == "RÉU" and l["2"]["corrigir_j"] == "AUTOR(A)"
    assert l["3"]["corrigir_j"] is None and l["3"]["j_corrigida"] is True
    assert tr.normalizar_j("Ré") == "RÉU" and tr.normalizar_j("AUTOR") == "AUTOR(A)" and tr.normalizar_j("x") is None
    assert tr.COL_PARALELA["corrigir_j"] == 26


# ----------------------------------------------------------------------------- fakes do Excel
class _Cells:
    def __init__(self, ws, r, c): self.ws, self.r, self.c = ws, r, c
    @property
    def Value(self): return self.ws.dados.get((self.r, self.c))
    @Value.setter
    def Value(self, v): self.ws.dados[(self.r, self.c)] = v; self.ws.eventos.append(("set", self.ws.Name, self.r, self.c, v))


class _Ws:
    def __init__(self, nome, eventos, procs, js):
        self.Name, self.eventos, self.dados, self.proxima = nome, eventos, {}, 3230
        for i, (p, j) in enumerate(zip(procs, js)):
            self.dados[(8 + i, 5)] = p; self.dados[(8 + i, 10)] = j
        self.n = len(procs)
    def Cells(self, r, c): return _Cells(self, r, c)


class _Wb:
    def __init__(self, abas): self.abas = abas
    def Worksheets(self, nome): return self.abas[nome]
    def Save(self): list(self.abas.values())[0].eventos.append(("Save",))


def _ambiente(monkeypatch, eventos):
    pub = _Ws("Publicações", eventos, ["0000001-11.2026.5.18.0001", "0000002-11.2026.5.18.0002", "0000001-11.2026.5.18.0001"], ["AUTOR(A)", "AUTOR(A)", "AUTOR(A)"])
    pf = _Ws("Prazos Feitos", eventos, ["0000001-11.2026.5.18.0001", "0000009-11.2026.5.18.0009"], ["AUTOR(A)", "RÉU"])
    wb = _Wb({"Publicações": pub, "Prazos Feitos": pf})
    monkeypatch.setattr(excel_com.time, "sleep", lambda s: None)
    monkeypatch.setattr(ap, "abrir_livro", lambda c, visivel=False: excel_com.Livro(object(), wb, False, c))
    monkeypatch.setattr(ap, "proxima_linha_tabela", lambda ws_, nome="Tabela1": (setattr(ws_, "proxima", ws_.proxima + 1), ws_.proxima)[1])
    monkeypatch.setattr(ap, "coluna_valores", lambda ws_, col, primeira=8: [ws_.dados.get((8 + i, col)) for i in range(ws_.n)])
    monkeypatch.setattr(ap, "set_cell", lambda ws_, r, c, v: setattr(ws_.Cells(r, c), "Value", v))
    monkeypatch.setattr(ap, "set_text", lambda ws_, r, c, v: setattr(ws_.Cells(r, c), "Value", v))
    monkeypatch.setattr(ap, "garantir_colunas_controle", lambda ws_, cfg_, **k: None)
    monkeypatch.setattr(ap, "avisar", lambda m: None); monkeypatch.setattr(ap, "_livro_aberto_no_excel", lambda c: None)
    return pub, pf


def test_correcao_de_j_em_todas_as_linhas_do_processo_e_nas_duas_abas(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); eventos = []
    pub, pf = _ambiente(monkeypatch, eventos)
    par = _paralela(cfg, date(2026, 10, 16), [_linha("1")], {"Z8": "RÉU", "M8": "Descartar"})     # só correção, sem Agendar
    razao = ap.transportar(cfg, par, str(tmp_path / "agenda.xlsx"))
    assert razao is not None
    assert pub.dados[(8, 10)] == "RÉU" and pub.dados[(10, 10)] == "RÉU" and pub.dados[(9, 10)] == "AUTOR(A)"
    assert pf.dados[(8, 10)] == "RÉU" and pf.dados[(9, 10)] == "RÉU"                                 # a 9 já era RÉU (outro processo)
    assert str(pub.dados[(8, 14)]).startswith("ADVOGADO ") and str(pf.dados[(8, 14)]).startswith("ADVOGADO ")
    assert (8, 14) not in pf.dados or pf.dados.get((9, 14)) is None                                   # outro processo não é tocado
    assert ("Save",) in eventos
    corr = [l for l in razao.linhas if l.get("tipo") == "j"]
    assert len(corr) == 3 and {(c["aba"], c["linha_agenda"]) for c in corr} == {("Publicações", 8), ("Publicações", 10), ("Prazos Feitos", 8)}
    assert all(c["j_antes"] == "AUTOR(A)" and c["j_depois"] == "RÉU" for c in corr)
    ws = openpyxl.load_workbook(par)["Lançamentos"]
    assert str(ws["Z8"].value).startswith("✔ RÉU")                                                    # marcada como aplicada


def test_correcao_ja_aplicada_nao_repete(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); eventos = []
    _ambiente(monkeypatch, eventos)
    par = _paralela(cfg, date(2026, 10, 16), [_linha("1")], {"Z8": "✔ RÉU (17/10 09:00)"})
    assert ap.transportar(cfg, par, str(tmp_path / "agenda.xlsx")) is None
    assert not any(e[0] == "set" for e in eventos)


def test_linha_transportada_leva_j_corrigida(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); eventos = []
    pub, _ = _ambiente(monkeypatch, eventos)
    par = _paralela(cfg, date(2026, 10, 16), [_linha("1")], {"Z8": "RÉU", "M8": "Agendar"})
    ap.transportar(cfg, par, str(tmp_path / "agenda.xlsx"))
    assert pub.dados[(3231, 10)] == "RÉU"                                                            # linha nova já com a J certa


def test_desfazer_restaura_j(cfg, tmp_path, monkeypatch):
    _cfg_pc(cfg, tmp_path); eventos = []
    pub, pf = _ambiente(monkeypatch, eventos)
    monkeypatch.setattr(ap, "_marcar_lote_na_paralela", lambda *a, **k: None)
    razao = tr.Razao(quando="x", paralela=str(tmp_path / "nao.xlsx"), tipo="transporte",
                     linhas=[{"tipo": "j", "aba": "Publicações", "linha_agenda": 8, "processo": "0000001-11.2026.5.18.0001",
                              "j_antes": "AUTOR(A)", "j_depois": "RÉU", "conf_antes": None}])
    pub.dados[(8, 10)] = "RÉU"; pub.dados[(8, 14)] = "ADVOGADO 17/10/2026"
    rz = str(tmp_path / "r.json"); razao.salvar(rz)
    ap.desfazer(cfg, rz)
    assert pub.dados[(8, 10)] == "AUTOR(A)" and pub.dados[(8, 14)] in (None, "")
