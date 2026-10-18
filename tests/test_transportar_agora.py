# -*- coding: utf-8 -*-
"""Gatilho manual do advogado: 'Transportar agora' grava na Agenda todas as decisões pendentes, sem esperar o vigia."""
import os
from datetime import date

import openpyxl

from paralela.planilha import Linha, gerar_paralela
from pc import agenda_paralela as ap
from pc.transporte import COL_PARALELA, Razao


def _cfg(cfg, tmp_path):
    cfg["pc"] = {"pasta_controladoria": str(tmp_path), "hora_alerta_sem_arquivo": "10:00", "intervalo_min": 15}
    cfg["subpasta_paralela"] = "CACT - Agenda Paralela"
    os.makedirs(tmp_path / "CACT - Agenda Paralela" / "razao", exist_ok=True)
    return cfg


def _linha(ids, processo):
    return Linha(prazo=date(2026, 10, 21), peca="Manifestação", publicado=date(2026, 10, 16), processo=processo, reclamante="A",
                 reclamado="R", advogado="E", prioridade="NORMAL", cliente="AUTOR(A)", despacho="r", destino="PRAZO DO ESCRITÓRIO", ids=ids)


def _paralela(cfg, dia, decisao):
    par = ap.caminho_paralela(cfg, dia)
    gerar_paralela([_linha("1" + dia.strftime("%d"), f"00000{dia.day:02d}-11.2026.5.18.0001")], [], par, dia, {"ok": True})
    if decisao:
        wb = openpyxl.load_workbook(par); wb["Lançamentos"].cell(8, COL_PARALELA["decisao"], decisao); wb.save(par)
    return par


def test_transportar_agora_cobre_hoje_e_dias_anteriores(cfg, tmp_path, monkeypatch):
    _cfg(cfg, tmp_path); avisos = []; monkeypatch.setattr(ap, "avisar", lambda m: avisos.append(m))
    p15 = _paralela(cfg, date(2026, 10, 15), "Agendar")
    _paralela(cfg, date(2026, 10, 16), "Descartar")                       # nada a transportar
    p18 = _paralela(cfg, date(2026, 10, 18), "Corrigir e agendar")        # hoje, ainda sem lote
    chamadas = []

    def transportar_falso(cfg_, caminho, agenda):
        chamadas.append(os.path.basename(caminho))
        return Razao(quando="x", paralela=caminho, linhas=[{"ids": ["1"]}, {"ids": ["2"]}] if caminho == p15 else [{"ids": ["3"]}])

    monkeypatch.setattr(ap, "transportar", transportar_falso)
    r = ap.transportar_agora(cfg, date(2026, 10, 18))
    assert chamadas == [os.path.basename(p15), os.path.basename(p18)]
    assert r == {"paralelas": 2, "linhas": 3, "erros": 0}
    assert avisos and "3 linha(s)" in avisos[-1]


def test_transportar_agora_sem_decisao_avisa_que_nao_ha_nada(cfg, tmp_path, monkeypatch):
    _cfg(cfg, tmp_path); avisos = []; monkeypatch.setattr(ap, "avisar", lambda m: avisos.append(m))
    _paralela(cfg, date(2026, 10, 18), None)
    monkeypatch.setattr(ap, "transportar", lambda *a: (_ for _ in ()).throw(AssertionError("não devia transportar")))
    assert ap.transportar_agora(cfg, date(2026, 10, 18)) == {"paralelas": 0, "linhas": 0, "erros": 0}
    assert avisos and "nenhuma" in avisos[-1].lower()


def test_transporte_exclusivo_entre_vigia_e_gatilho_manual(cfg, tmp_path, monkeypatch):
    """Vigia e 'Transportar agora' não gravam ao mesmo tempo: trava pc/transporte.lock; quem não a obtém no prazo desiste e avisa."""
    _cfg(cfg, tmp_path); avisos = []; monkeypatch.setattr(ap, "avisar", lambda m: avisos.append(m))
    monkeypatch.setattr(ap, "TRAVA_TRANSPORTE", str(tmp_path / "transporte.lock"))
    monkeypatch.setattr(ap, "transportar", lambda *a: (_ for _ in ()).throw(AssertionError("não devia transportar com a trava ocupada")))
    fh = ap.obter_trava(str(tmp_path / "transporte.lock"))
    try:
        ap._transportar_seguro(cfg, "x.xlsx", espera_s=0)
    finally:
        ap.soltar_trava(fh)
    assert avisos and "outro transporte" in avisos[-1].lower()
    chamadas = []
    monkeypatch.setattr(ap, "transportar", lambda cfg_, c, a: chamadas.append(c) or None)
    ap._transportar_seguro(cfg, "x.xlsx", espera_s=0)                      # trava livre: transporta
    assert chamadas == ["x.xlsx"]


def test_cli_agora(cfg, tmp_path, monkeypatch):
    monkeypatch.setattr(ap, "carregar_cfg", lambda: cfg)
    chamadas = []
    monkeypatch.setattr(ap, "transportar_agora", lambda cfg_, hoje: chamadas.append(hoje) or {"paralelas": 0, "linhas": 0, "erros": 0})
    assert ap.main(["agora"]) == 0 and chamadas == [date.today()]


def _pronto(monkeypatch):
    monkeypatch.setattr(ap, "arquivo_pronto", lambda c, agora=None: True)
    monkeypatch.setattr(ap, "abrir_para_o_advogado", lambda c, n=0: None)
    monkeypatch.setattr(ap, "coletar", lambda *a, **k: {"novos": 0, "registros": 0, "caminho": None})
    monkeypatch.setattr(ap, "verificar_persistencia", lambda *a, **k: [])


def test_vigia_nao_transporta_quando_transporte_automatico_desligado(cfg, tmp_path, monkeypatch):
    """Decisão do advogado (18/10/2026): a Agenda só recebe as linhas quando ele dispara 'Transportar agora'."""
    _cfg(cfg, tmp_path); cfg["pc"]["transporte_automatico"] = False
    avisos = []; monkeypatch.setattr(ap, "avisar", lambda m: avisos.append(m)); _pronto(monkeypatch)
    monkeypatch.setattr(ap, "_transportar_seguro", lambda *a, **k: (_ for _ in ()).throw(AssertionError("não devia transportar sozinho")))
    _paralela(cfg, date(2026, 10, 15), "Agendar")                       # dia anterior com decisão pendente
    hoje = _paralela(cfg, date(2026, 10, 18), None)
    estado = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 18), date(2026, 10, 18) and __import__("datetime").datetime(2026, 10, 18, 9, 0), estado)   # abre
    wb = openpyxl.load_workbook(hoje); wb["Lançamentos"].cell(8, COL_PARALELA["decisao"], "Agendar"); wb.save(hoje)      # advogado marca
    os.utime(hoje, None)
    ap.passo_vigia(cfg, date(2026, 10, 18), __import__("datetime").datetime(2026, 10, 18, 9, 15), estado)
    assert not any("gravada" in a for a in avisos)


def test_vigia_avisa_decisoes_pendentes_no_horario_configurado(cfg, tmp_path, monkeypatch):
    from datetime import datetime
    _cfg(cfg, tmp_path); cfg["pc"]["transporte_automatico"] = False; cfg["pc"]["hora_alerta_transporte"] = "17:30"
    avisos = []; monkeypatch.setattr(ap, "avisar", lambda m: avisos.append(m)); _pronto(monkeypatch)
    _paralela(cfg, date(2026, 10, 15), "Agendar"); _paralela(cfg, date(2026, 10, 18), "Corrigir e agendar")
    estado = ap.Estado()
    ap.passo_vigia(cfg, date(2026, 10, 18), datetime(2026, 10, 18, 17, 0), estado)
    assert not any("Transportar agora" in a for a in avisos)
    ap.passo_vigia(cfg, date(2026, 10, 18), datetime(2026, 10, 18, 17, 31), estado)
    ap.passo_vigia(cfg, date(2026, 10, 18), datetime(2026, 10, 18, 17, 46), estado)
    pend = [a for a in avisos if "Transportar agora" in a]
    assert len(pend) == 1 and "2 decis" in pend[0]


def test_vigia_transporta_sozinho_quando_automatico_ligado(cfg, tmp_path, monkeypatch):
    from datetime import datetime
    _cfg(cfg, tmp_path)                                                   # padrão: automático
    monkeypatch.setattr(ap, "avisar", lambda m: None); _pronto(monkeypatch)
    chamadas = []; monkeypatch.setattr(ap, "_transportar_seguro", lambda cfg_, c, **k: chamadas.append(os.path.basename(c)))
    p15 = _paralela(cfg, date(2026, 10, 15), "Agendar"); _paralela(cfg, date(2026, 10, 18), None)
    ap.passo_vigia(cfg, date(2026, 10, 18), datetime(2026, 10, 18, 9, 0), ap.Estado())
    assert chamadas == [os.path.basename(p15)]
