# -*- coding: utf-8 -*-
"""Conferência PUSH: cruza as intimações recebidas por e-mail (relatório `Relatorio_PUSH_AAAA-MM-DD_0900.xlsx`,
gerado às 9h a partir do Gmail: PJe Push, TRT-Push, Jusbrasil) com a Paralela (hoje e ontem) e a Agenda.

Intimação do Push sem cobertura = intimação feita pelo sistema do PJe, sem publicação no DJEN: vira linha na aba
Lançamentos (PRAZO DO ESCRITÓRIO + CONFERIR) para o advogado decidir. Os demais eventos (documento sigiloso,
juntadas, pautas) ficam só na aba "Conferência PUSH", para conhecimento.
"""
from __future__ import annotations

import copy
import glob
import hashlib
import os
import re
import shutil
import tempfile
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from pc.transporte import COL_PARALELA, PRIMEIRA_LINHA

CNJ = re.compile(r"\d{7}-?\d{2}\.?\d{4}\.?\d\.?\d{2}\.?\d{4}")
ABA_CONFERENCIA = "Conferência PUSH"
JANELA_AGENDA_DIAS = 5
PRAZO_PADRAO = 5
TIPOS_INTIMACAO = ("INTIMA", "PUBLICA", "CITA")
SITUACOES = ("SEM COBERTURA", "coberto pela Paralela", "coberto pela Agenda", "informativo")
CAB_CONF = ["Data", "Processo", "Tribunal", "Tipo de andamento", "Situação", "Detalhe", "Teor (início)"]
MOTIVO_PUSH = "intimação pelo PJe (Push) sem publicação no DJEN: conferir teor e prazo no PJe"


def _norm(v) -> str | None:
    m = CNJ.search(str(v or ""))
    return re.sub(r"\D", "", m.group()) if m else None


def _mascara(v) -> str:
    m = CNJ.search(str(v or ""))
    if not m:
        return str(v or "")
    d = re.sub(r"\D", "", m.group())
    return f"{d[:7]}-{d[7:9]}.{d[9:13]}.{d[13]}.{d[14:16]}.{d[16:]}"


def _data(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.fromisoformat(str(v)[:19]).date()
    except (ValueError, TypeError):
        return None


def _copia(caminho: str) -> str:
    tmp = os.path.join(tempfile.mkdtemp(prefix="push_"), os.path.basename(caminho))
    shutil.copyfile(caminho, tmp)
    return tmp


# ----------------------------------------------------------------------------- caminhos
def pasta_push(cfg: dict) -> str:
    return os.path.join((cfg.get("pc") or {}).get("pasta_controladoria", ""), cfg.get("subpasta_push", "CACT - PUSH Diário"))


def caminho_push(cfg: dict, dia: date) -> str | None:
    exato = os.path.join(pasta_push(cfg), f"Relatorio_PUSH_{dia:%Y-%m-%d}_0900.xlsx")
    if os.path.exists(exato):
        return exato
    outros = sorted(glob.glob(os.path.join(pasta_push(cfg), f"Relatorio_PUSH_{dia:%Y-%m-%d}_*.xlsx")))
    return outros[-1] if outros else None


# ----------------------------------------------------------------------------- leitura
def ler_push(caminho: str) -> list[dict]:
    tmp = _copia(caminho)
    wb = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
    out = []
    try:
        ws = wb["Publicações e Andamentos"] if "Publicações e Andamentos" in wb.sheetnames else wb.worksheets[0]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 9:
                continue
            norm = _norm(row[2])
            if not norm:
                continue
            tipo = str(row[8] or "").strip()
            teor = str(row[9] or "").strip() if len(row) > 9 else ""
            data_ev = _data(row[0])
            ident = "PUSH:" + hashlib.sha1(f"{norm}|{data_ev}|{tipo}|{teor[:60]}".encode("utf-8")).hexdigest()[:12]
            out.append({"data": data_ev, "norm": norm, "processo": _mascara(row[2]), "tribunal": str(row[3] or ""), "tipo": tipo,
                        "teor": teor, "intimacao": any(t in tipo.upper() for t in TIPOS_INTIMACAO), "ident": ident})
    finally:
        wb.close()
        shutil.rmtree(os.path.dirname(tmp), ignore_errors=True)
    return out


def _procs_paralela(caminho: str) -> set[str]:
    tmp = _copia(caminho)
    wb = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
    out: set[str] = set()
    try:
        for aba in ("Lançamentos", "Fora da Agenda"):
            if aba in wb.sheetnames:
                for row in wb[aba].iter_rows(min_row=PRIMEIRA_LINHA, values_only=True):
                    n = _norm(row[4]) if row and len(row) > 4 else None
                    if n:
                        out.add(n)
    finally:
        wb.close()
        shutil.rmtree(os.path.dirname(tmp), ignore_errors=True)
    return out


def _agenda(caminho: str) -> tuple[dict[str, list[tuple[date, date]]], dict[str, dict]]:
    """norm → [(publicado, prazo)] e norm → {reclamante, reclamado, advogado, cliente} (última linha vista)."""
    tmp = _copia(caminho)
    wb = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
    datas: dict[str, list] = {}
    ref: dict[str, dict] = {}
    try:
        for aba in ("Publicações", "Prazos Feitos"):
            if aba not in wb.sheetnames:
                continue
            for row in wb[aba].iter_rows(min_row=PRIMEIRA_LINHA, values_only=True):
                n = _norm(row[4]) if row and len(row) > 4 else None
                if not n:
                    continue
                pub, prazo = _data(row[3]), _data(row[0])
                if pub:
                    datas.setdefault(n, []).append((pub, prazo))
                if len(row) > 9:
                    ref[n] = {"reclamante": row[5], "reclamado": row[6], "advogado": row[7], "cliente": row[9]}
    finally:
        wb.close()
        shutil.rmtree(os.path.dirname(tmp), ignore_errors=True)
    return datas, ref


# ----------------------------------------------------------------------------- classificação
def classificar(eventos: list[dict], paralelas: dict[str, date], agenda: dict[str, list], dia: date) -> list[dict]:
    out = []
    for e in eventos:
        e = dict(e)
        if not e.get("intimacao"):
            e["situacao"], e["detalhe"] = "informativo", ""
        elif e["norm"] in paralelas:
            e["situacao"], e["detalhe"] = "coberto pela Paralela", f"Paralela de {paralelas[e['norm']]:%d/%m}"
        else:
            recentes = [(p, z) for p, z in agenda.get(e["norm"], []) if p and p >= dia - timedelta(days=JANELA_AGENDA_DIAS)]
            if recentes:
                p, z = max(recentes)
                e["situacao"], e["detalhe"] = "coberto pela Agenda", f"publicado {p:%d/%m}, prazo {z:%d/%m}" if z else f"publicado {p:%d/%m}"
            else:
                e["situacao"], e["detalhe"] = "SEM COBERTURA", "intimação pelo PJe sem publicação no DJEN"
        out.append(e)
    return out


# ----------------------------------------------------------------------------- escrita na Paralela
def _dias_uteis(cfg: dict, base: date, n: int, tribunal: str) -> date:
    from paralela.prazos import carregar_calendario, dias_uteis_apos, suspensoes
    feriados, _ = carregar_calendario(cfg, tribunal or "TRT18")
    return dias_uteis_apos(base, n, feriados, suspensoes(cfg, tribunal or "TRT18"))


def _livro_aberto_no_excel(caminho: str):
    from pc.excel_com import _livro_na_instancia_ativa
    return _livro_na_instancia_ativa(caminho)


def _linha_lancamento(cfg: dict, e: dict, ref: dict, dia: date) -> dict[int, object]:
    base = e.get("data") or dia
    prazo = _dias_uteis(cfg, base, PRAZO_PADRAO, e.get("tribunal", ""))
    r = ref.get(e["norm"], {})
    teor = re.sub(r"\s+", " ", e.get("teor") or "")[:250]
    return {1: prazo, 3: "Verificar intimação no PJe (Push)", 4: base, 5: e["processo"], 6: r.get("reclamante"), 7: r.get("reclamado"),
            8: r.get("advogado"), 9: "NORMAL", 10: r.get("cliente"), 11: f"PUSH: {e.get('tipo')} · {teor}",
            COL_PARALELA["destino"]: "PRAZO DO ESCRITÓRIO", COL_PARALELA["motivo"]: MOTIVO_PUSH, 15: "?", 17: "PUSH (e-mail)",
            18: None, COL_PARALELA["ids"]: e["ident"], COL_PARALELA["prazo_calculado"]: prazo}


def _ids_existentes(ws) -> set[str]:
    return {str(ws.cell(r, COL_PARALELA["ids"]).value) for r in range(PRIMEIRA_LINHA, ws.max_row + 1) if ws.cell(r, COL_PARALELA["ids"]).value}


def _ultima_linha_dados(ws) -> int:
    ultima = PRIMEIRA_LINHA - 1
    for r in range(PRIMEIRA_LINHA, ws.max_row + 1):
        if ws.cell(r, 5).value:
            ultima = r
    return ultima


def _inserir_linha_openpyxl(ws, valores: dict[int, object]) -> int:
    r = _ultima_linha_dados(ws) + 1
    ws.insert_rows(r)
    modelo = r - 1 if r - 1 >= PRIMEIRA_LINHA else None
    for c in range(1, 27):
        cel = ws.cell(r, c)
        if modelo:
            m = ws.cell(modelo, c)
            cel.font, cel.border, cel.alignment, cel.number_format = copy.copy(m.font), copy.copy(m.border), copy.copy(m.alignment), m.number_format
        cel.fill = PatternFill("solid", fgColor="FCE4D6")
    for c, v in valores.items():
        ws.cell(r, c, v)
    ws.cell(r, 2, f'=IF(A{r}="","",A{r})')
    ws.cell(r, 13).fill = PatternFill("solid", fgColor="FFFF00"); ws.cell(r, 26).fill = PatternFill("solid", fgColor="FFFF00")
    for dv in ws.data_validations.dataValidation:
        s = str(dv.sqref)
        if s.startswith("M"):
            dv.add(f"M{r}")
        elif s.startswith("Z"):
            dv.add(f"Z{r}")
    return r


def _escrever_conferencia_openpyxl(wb, classificados: list[dict], dia: date, fonte: str) -> None:
    if ABA_CONFERENCIA in wb.sheetnames:
        del wb[ABA_CONFERENCIA]
    ws = wb.create_sheet(ABA_CONFERENCIA)
    ws["A1"] = f"Conferência PUSH de {dia:%d/%m/%Y} — intimações recebidas por e-mail x Paralela x Agenda"
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"] = (f"Fonte: {os.path.basename(fonte)}. Intimação SEM COBERTURA = feita pelo sistema do PJe, sem publicação no DJEN: "
                f"virou linha na aba Lançamentos (laranja, CONFERIR). Os demais eventos são só para conhecimento.")
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="595959")
    thin = Side(style="thin", color="BFBFBF"); borda = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(CAB_CONF, 1):
        cel = ws.cell(7, c, h); cel.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="0070C0"); cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cel.border = borda
    for w, col in zip([12, 26, 9, 30, 20, 30, 70], "ABCDEFG"):
        ws.column_dimensions[col].width = w
    cor = {"SEM COBERTURA": "FCE4D6", "coberto pela Paralela": "E2EFDA", "coberto pela Agenda": "E2EFDA", "informativo": "F2F2F2"}
    ordem = {s: i for i, s in enumerate(SITUACOES)}
    for i, e in enumerate(sorted(classificados, key=lambda x: (ordem.get(x["situacao"], 9), x.get("data") or dia)), 8):
        vals = [e.get("data"), e["processo"], e.get("tribunal"), e.get("tipo"), e["situacao"], e.get("detalhe"), (e.get("teor") or "")[:200]]
        for c, v in enumerate(vals, 1):
            cel = ws.cell(i, c, v); cel.font = Font(name="Courier New" if c == 2 else "Arial", size=9)
            cel.border = borda; cel.fill = PatternFill("solid", fgColor=cor.get(e["situacao"], "FFFFFF"))
            cel.alignment = Alignment(vertical="top", wrap_text=(c in (4, 6, 7)))
            if c == 1:
                cel.number_format = "dd/mm/yy;@"
    ws.freeze_panes = "A8"


def conferir_push(cfg: dict, dia: date) -> list[dict]:
    """Lê o relatório PUSH do dia, classifica cada evento e grava a aba 'Conferência PUSH' na Paralela de hoje;
    intimações SEM COBERTURA viram linhas em Lançamentos (idempotente pelo ID 'PUSH:…')."""
    from pc.agenda_paralela import _log, caminho_agenda, caminho_paralela
    fonte = caminho_push(cfg, dia)
    paralela_hoje = caminho_paralela(cfg, dia)
    if not fonte or not os.path.exists(paralela_hoje):
        return []
    eventos = ler_push(fonte)
    paralelas: dict[str, date] = {}
    for d in (dia - timedelta(days=1), dia):
        c = caminho_paralela(cfg, d)
        if os.path.exists(c):
            for n in _procs_paralela(c):
                paralelas[n] = d
    agenda_datas, ref = _agenda(caminho_agenda(cfg)) if os.path.exists(caminho_agenda(cfg)) else ({}, {})
    classificados = classificar(eventos, paralelas, agenda_datas, dia)
    sem_cobertura = [e for e in classificados if e["situacao"] == "SEM COBERTURA"]

    aberto = _livro_aberto_no_excel(paralela_hoje)
    if aberto:
        # Paralela aberta no Excel do advogado: gravar pela instância (aba de conferência + linhas novas)
        from pc.excel_com import com, set_cell, set_text
        _, wb = aberto
        nomes = [com(lambda i=i: wb.Worksheets(i).Name) for i in range(1, int(com(lambda: wb.Worksheets.Count)) + 1)]
        if ABA_CONFERENCIA in nomes:
            ws_c = com(wb.Worksheets, ABA_CONFERENCIA)
            com(lambda: ws_c.Cells.Clear())
        else:
            ws_c = com(lambda: wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count)))
            com(lambda: setattr(ws_c, "Name", ABA_CONFERENCIA))
        com(set_cell, ws_c, 1, 1, f"Conferência PUSH de {dia:%d/%m/%Y} — intimações recebidas por e-mail x Paralela x Agenda")
        for c, h in enumerate(CAB_CONF, 1):
            com(set_cell, ws_c, 7, c, h)
        ordem = {s: i for i, s in enumerate(SITUACOES)}
        for i, e in enumerate(sorted(classificados, key=lambda x: (ordem.get(x["situacao"], 9), x.get("data") or dia)), 8):
            for c, v in enumerate([e.get("data"), e["processo"], e.get("tribunal"), e.get("tipo"), e["situacao"], e.get("detalhe"), (e.get("teor") or "")[:200]], 1):
                if v is not None:
                    com(set_cell, ws_c, i, c, v if not isinstance(v, date) else datetime(v.year, v.month, v.day))
        ws_l = com(wb.Worksheets, "Lançamentos")
        ultima = int(com(lambda: ws_l.Cells(ws_l.Rows.Count, 5).End(-4162).Row))
        existentes = {str(com(lambda r=r: ws_l.Cells(r, COL_PARALELA["ids"]).Value) or "") for r in range(PRIMEIRA_LINHA, ultima + 1)}
        novas = 0
        for e in sem_cobertura:
            if e["ident"] in existentes:
                continue
            r = ultima + 1
            com(lambda r=r: ws_l.Rows(r).Insert())
            for c, v in _linha_lancamento(cfg, e, ref, dia).items():
                if v is None:
                    continue
                if c in (5, COL_PARALELA["ids"]):
                    com(set_text, ws_l, r, c, v)
                else:
                    com(set_cell, ws_l, r, c, v if not isinstance(v, date) else datetime(v.year, v.month, v.day))
            com(set_cell, ws_l, r, 2, f'=IF(A{r}="","",A{r})')
            ultima += 1; novas += 1
        com(wb.Save)
    else:
        wb = openpyxl.load_workbook(paralela_hoje)
        _escrever_conferencia_openpyxl(wb, classificados, dia, fonte)
        ws_l = wb["Lançamentos"]
        existentes = _ids_existentes(ws_l)
        novas = 0
        for e in sem_cobertura:
            if e["ident"] in existentes:
                continue
            _inserir_linha_openpyxl(ws_l, _linha_lancamento(cfg, e, ref, dia)); novas += 1
        wb.save(paralela_hoje)
    _log(f"Conferência PUSH {dia:%d/%m}: {len(eventos)} eventos, {sum(1 for e in classificados if e['intimacao'])} intimações, "
         f"{len(sem_cobertura)} sem cobertura, {novas} linha(s) nova(s) em Lançamentos")
    return classificados
