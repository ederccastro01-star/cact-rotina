# -*- coding: utf-8 -*-
"""Operação assistida (spec §8): compara a Paralela com os lançamentos MANUAIS da Agenda no mesmo período.

Critério de saída da fase assistida: nenhuma data automática POSTERIOR à manual.
Uso: python -m pc.comparar_manual [--inicio DD.MM.AAAA] [--fim DD.MM.AAAA]
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import sys
import tempfile
from datetime import date, datetime, timedelta

import openpyxl

from pc.transporte import COL, ler_paralela

CNJ = re.compile(r"\d{7}-?\d{2}\.?\d{4}\.?\d\.?\d{2}\.?\d{4}")
JANELA_DIAS = 4


def _norm(v) -> str | None:
    m = CNJ.search(str(v or ""))
    return re.sub(r"\D", "", m.group()) if m else None


def _data(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _copia(caminho: str) -> str:
    from pc import arquivos
    return arquivos.copia_temporaria(caminho, prefixo="cmp_")      # CopyFileW: a Agenda fica trancada pelo AutoSave


def linhas_manuais(agenda: str, inicio: date, fim: date, abas: tuple[str, ...] = ("Publicações", "Prazos Feitos")) -> list[dict]:
    """Linhas da Agenda (Publicações E Prazos Feitos) com Publicado no intervalo e Origem (Q) vazia (= à mão)."""
    tmp = _copia(agenda)
    wb = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
    out = []
    try:
        for aba in abas:
          if aba not in wb.sheetnames:
            continue
          for row in wb[aba].iter_rows(min_row=8, values_only=True):
            if not row or len(row) < 5:
                continue
            pub = _data(row[3]); proc = _norm(row[4])
            origem = row[COL["origem"] - 1] if len(row) >= COL["origem"] else None
            if not proc or not pub or origem not in (None, "") or not (inicio <= pub <= fim):
                continue
            out.append({"proc": proc, "processo": str(row[4]).strip(), "prazo": _data(row[0]), "peca": str(row[2] or ""), "publicado": pub})
    finally:
        wb.close()
        shutil.rmtree(os.path.dirname(tmp), ignore_errors=True)
    return out


def linhas_paralela(paralelas: list[str]) -> list[dict]:
    out = []
    for c in paralelas:
        for l in ler_paralela(c):
            if l["aba"] != "Lançamentos":
                continue
            a = l["a_k"]
            prazo_maquina = _data(l.get("prazo_calculado")) or _data(a[0])     # Y: a data da máquina, não a corrigida pelo advogado
            out.append({"proc": _norm(a[4]), "processo": str(a[4]), "prazo": prazo_maquina, "peca": str(a[2] or ""),
                        "publicado": _data(a[3]), "destino": l.get("destino"), "arquivo": os.path.basename(c)})
    return out


def comparar(agenda: str, paralelas: list[str], inicio: date, fim: date) -> dict:
    manuais = linhas_manuais(agenda, inicio, fim)
    autos = [p for p in linhas_paralela(paralelas) if p["publicado"] and inicio <= p["publicado"] <= fim]
    casadas, posterior, usados = 0, [], set()
    for m in manuais:
        cand = [(i, p) for i, p in enumerate(autos) if i not in usados and p["proc"] == m["proc"] and p["publicado"]
                and abs((p["publicado"] - m["publicado"]).days) <= JANELA_DIAS]
        if not cand:
            continue
        # a linha automática mais próxima em data de prazo
        i, p = min(cand, key=lambda ip: abs(((ip[1]["prazo"] or date.max) - (m["prazo"] or date.max)).days))
        usados.add(i); casadas += 1
        if m["prazo"] and p["prazo"] and p["prazo"] > m["prazo"]:
            posterior.append((m["processo"], p["prazo"], m["prazo"]))
    return {"manuais": len(manuais), "casadas": casadas, "data_auto_posterior": posterior,
            "so_na_paralela": len(autos) - len(usados), "so_na_agenda": len(manuais) - casadas}


def relatorio(r: dict, inicio: date, fim: date) -> str:
    linhas = [f"Comparação Paralela x lançamentos manuais — publicado de {inicio:%d/%m/%Y} a {fim:%d/%m/%Y}",
              f"manuais: {r['manuais']} | casadas com a Paralela: {r['casadas']} | só na Agenda: {r['so_na_agenda']} | só na Paralela: {r['so_na_paralela']}",
              f"DATA AUTOMÁTICA POSTERIOR À MANUAL: {len(r['data_auto_posterior'])}"]
    for proc, auto, manual in r["data_auto_posterior"]:
        linhas.append(f"  {proc}: Paralela {auto:%d/%m/%Y} x manual {manual:%d/%m/%Y}")
    linhas.append("Critério de saída (spec §8): zero casos de data automática posterior." if not r["data_auto_posterior"]
                  else "Critério de saída NÃO atendido: corrigir antes de confiar no transporte.")
    return "\n".join(linhas)


def main(argv=None) -> int:
    from pc.agenda_paralela import carregar_cfg, caminho_agenda, pasta_paralela
    ap = argparse.ArgumentParser(description="Paralela x lançamentos manuais")
    ap.add_argument("--inicio"); ap.add_argument("--fim")
    args = ap.parse_args(argv)
    cfg = carregar_cfg()
    fim = datetime.strptime(args.fim, "%d.%m.%Y").date() if args.fim else date.today()
    inicio = datetime.strptime(args.inicio, "%d.%m.%Y").date() if args.inicio else fim - timedelta(days=7)
    pasta = pasta_paralela(cfg)
    paralelas = sorted(os.path.join(pasta, f) for f in os.listdir(pasta) if f.startswith("CACT - Agenda Paralela") and f.endswith(".xlsx"))
    r = comparar(caminho_agenda(cfg), paralelas, inicio, fim)
    txt = relatorio(r, inicio, fim)
    saida = os.path.join(pasta, f"comparacao-{fim:%Y-%m-%d}.txt")
    with open(saida, "w", encoding="utf-8") as fh:
        fh.write(txt + "\n")
    print(txt); print("->", saida)
    return 0


if __name__ == "__main__":
    sys.exit(main())
