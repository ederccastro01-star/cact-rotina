# -*- coding: utf-8 -*-
"""Gera `Relatorio_PUSH_AAAA-MM-DD_HHMM.xlsx` no formato da antiga rotina de nuvem (spec §9c): aba "Publicações e
Andamentos" com 18 colunas, uma linha por evento/publicação, e aba "Resumo" com fórmulas (nunca valor calculado).
`pc/push.py` consome a aba principal pelas colunas 0 (data), 2 (processo), 3 (tribunal), 8 (tipo) e 9 (teor).
"""
from __future__ import annotations

import hashlib
import os
from collections import Counter
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pc.push_fontes import Registro

ABA = "Publicações e Andamentos"
COLUNAS = ["Data do andamento", "Data de autuação", "Nº do processo (CNJ)", "Tribunal", "Ramo da Justiça", "Grau", "Classe judicial",
           "Órgão / Vara", "Tipo de andamento", "Andamento / Teor", "Autor", "Réu", "Origem (sistema)", "Data do e-mail",
           "Assunto do e-mail", "Remetente", "ID da conversa", "Registro repetido?"]
LARGURAS = [18, 18, 26, 10, 20, 10, 30, 44, 30, 90, 32, 32, 18, 18, 46, 30, 20, 20]
COLS_WRAP = ("H", "J", "K", "L", "O")
COLS_DATA = ("A", "B", "N")
FORMATO_DATA = "DD/MM/YYYY HH:MM"
FONTE = "Arial"
NOTAS = [
    "1. Um mesmo e-mail PUSH traz vários eventos; cada um vira uma linha.",
    "2. O PUSH do PJe reenvia o processo a cada movimentação; repetições entre e-mails de dias diferentes são esperadas e não são erro.",
    "3. Teores do Jusbrasil podem chegar truncados (terminam em \"...\"); texto integral só pelo link original.",
    "4. E-mails do PUSH legado chegam com acentuação corrompida na origem.",
]


def id_conversa(message_id: str) -> str:
    return hashlib.sha1((message_id or "").encode("utf-8")).hexdigest()[:16]


def _ordem(r: Registro):
    return (r.data or datetime.min, r.processo)


def gerar_relatorio(registros: list[Registro], caminho: str, janela: tuple[datetime, datetime], contas: list[str],
                    backlog: int = 0, gerado_em: datetime | None = None, faltantes_jusbrasil: int = 0) -> dict:
    gerado_em = gerado_em or datetime.now()
    conversas = len({r.message_id for r in registros})
    info = {"registros": len(registros), "conversas": conversas, "caminho": None, "backlog": backlog}
    if not registros:
        return info
    regs = sorted(registros, key=_ordem, reverse=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ABA
    _aba_principal(ws, regs)
    _aba_resumo(wb.create_sheet("Resumo"), regs, janela, contas, backlog, gerado_em, conversas, faltantes_jusbrasil)
    os.makedirs(os.path.dirname(caminho) or ".", exist_ok=True)
    wb.save(caminho)
    info["caminho"] = caminho
    return info


def _aba_principal(ws, regs: list[Registro]) -> None:
    cab_font = Font(name=FONTE, bold=True, color="FFFFFF")
    cab_fill = PatternFill("solid", fgColor="1F3864")
    for c, nome in enumerate(COLUNAS, 1):
        cel = ws.cell(1, c, nome)
        cel.font = cab_font
        cel.fill = cab_fill
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    vistos: Counter = Counter()
    par = PatternFill("solid", fgColor="F2F2F2")
    for i, r in enumerate(regs, 2):
        chave = (r.processo, r.data, r.teor)
        vistos[chave] += 1
        rep = "Não" if vistos[chave] == 1 else f"Sim ({vistos[chave]}ª ocorrência)"
        valores = [r.data, r.autuacao, r.processo, r.tribunal, r.ramo, r.grau, r.classe, r.orgao, r.tipo, r.teor, r.autor, r.reu,
                   r.origem, r.data_email, r.assunto, r.remetente, id_conversa(r.message_id), rep]
        for c, v in enumerate(valores, 1):
            cel = ws.cell(i, c, v if v != "" else None)
            cel.font = Font(name=FONTE)
            letra = get_column_letter(c)
            if letra in COLS_DATA:
                cel.number_format = FORMATO_DATA
            cel.alignment = Alignment(vertical="top", wrap_text=letra in COLS_WRAP)
            if (i - 2) % 2 == 1:
                cel.fill = par
    for c, w in enumerate(LARGURAS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:R{len(regs) + 1}"


def _aba_resumo(rs, regs: list[Registro], janela, contas, backlog, gerado_em, conversas, faltantes) -> None:
    n = len(regs) + 1
    ref = f"'{ABA}'!"
    linhas: list[list] = [
        ["Relatório de Publicações e Andamentos Processuais"],
        [f"Fonte: e-mails lidos no Outlook (contas: {', '.join(contas)}): PJe Push, TRT-Push legado, Jusbrasil e avisos de tribunais"],
        [f"Janela coberta: {janela[0]:%d/%m/%Y %H:%M} a {janela[1]:%d/%m/%Y %H:%M} (e-mails ainda não processados nesse período)"],
        [f"Data de geração: {gerado_em:%d/%m/%Y %H:%M}"],
        ["Total de registros", f"=COUNTA({ref}C2:C{n})"],
        ["Total de conversas processadas", conversas],
        ["Backlog (>janela) deixado fora do relatório", backlog],
        [],
        ["NOTAS DE QUALIDADE"],
    ]
    for nota in NOTAS:
        linhas.append([nota])
    linhas.append([f"5. Backlog (e-mails elegíveis mais antigos que a janela) deixado fora do relatório: {backlog}."])
    if faltantes:
        linhas.append([f"6. O digest do Jusbrasil 'no seu nome' (Hotmail) não trouxe {faltantes} publicação(ões): só no site do Jusbrasil."])
    linhas.append([])
    total_row = 5
    blocos = [("POR TRIBUNAL", "D", lambda r: r.tribunal), ("POR RAMO DA JUSTIÇA", "E", lambda r: r.ramo),
              ("POR TIPO DE ANDAMENTO", "I", lambda r: r.tipo), ("POR ORIGEM (SISTEMA)", "M", lambda r: r.origem)]
    titulos = set()
    for titulo, col, chave in blocos:
        linhas.append([titulo, "Qtd", "%"])
        titulos.add(len(linhas))
        inicio = len(linhas) + 1
        for valor, _ in Counter(chave(r) or "(vazio)" for r in regs).most_common():
            r_ = len(linhas) + 1
            linhas.append([valor, f"=COUNTIF({ref}${col}$2:${col}${n},A{r_})", f"=B{r_}/$B${total_row}"])
        fim = len(linhas)
        linhas.append(["Total", f"=SUM(B{inicio}:B{fim})", f"=SUM(C{inicio}:C{fim})"])
        linhas.append([])
    linhas.append(["TOP 20 PROCESSOS COM MAIS REGISTROS", "Qtd"])
    titulos.add(len(linhas))
    for proc, _ in Counter(r.processo for r in regs if r.processo).most_common(20):
        r_ = len(linhas) + 1
        linhas.append([proc, f"=COUNTIF({ref}$C$2:$C${n},A{r_})"])
    negrito = Font(name=FONTE, bold=True)
    for i, linha in enumerate(linhas, 1):
        for c, v in enumerate(linha, 1):
            cel = rs.cell(i, c, v)
            cel.font = negrito if (i == 1 or i in titulos or (linha and linha[0] in ("Total", "NOTAS DE QUALIDADE"))) else Font(name=FONTE)
            if c == 3 and isinstance(v, str) and v.startswith("=") and "SUM" not in v:
                cel.number_format = "0.0%"
            elif c == 3 and isinstance(v, str) and v.startswith("=SUM"):
                cel.number_format = "0.0%"
    rs.column_dimensions["A"].width = 60
    rs.column_dimensions["B"].width = 14
    rs.column_dimensions["C"].width = 10
